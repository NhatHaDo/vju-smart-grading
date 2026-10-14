"""
Trình bày đề khi trộn (2026-10-14), like YoungMix's options:

  - header  "none"   = as before (tên bộ đề, "Mã đề: …");
            "vju"    = đầu đề của Trường ĐH Việt Nhật (ĐHQGHN / trường / khoa
                       on the left, "ĐỀ THI KẾT THÚC HỌC PHẦN" / học kỳ on the
                       right, then học phần and thời gian centered);
            "school" = đầu đề phổ thông (Sở / trường on the left, tên bài
                       kiểm tra / năm học / môn / thời gian on the right);
  - label   "." or ":"  → "Câu 1." / "Câu 1:";
  - start   the number of the first câu of every part (đề phần 2 bắt đầu ở câu 41…);
  - english "Question 1.", "PART", "Exam code", "THE END"…

Applied when the Word files are made (exam_mixer.version_docx, the
format-keeping docx_mixer and the đáp án), so it can be changed after mixing
without changing the mã đề. Kept in the bộ đề's settings["format"].
"""
from __future__ import annotations

from app.services.question_io import PART_NAME

HEADERS = ("none", "vju", "school")
# what the teacher types in "Trình bày đề"; empty = the default below
FIELDS = ("org1", "org2", "org3", "title", "term", "subject", "subject_code", "duration")
MAX_START = 999

_VN = {
    "q": "Câu", "code": "Mã đề", "part": "PHẦN", "end": "— HẾT —", "parts": PART_NAME,
    "name": "Họ và tên: ……………………………………………   Số báo danh: ……………………",
    "course": "Học phần", "course_code": "Mã học phần", "subject": "Môn", "time": "Thời gian làm bài",
    "minutes": "phút", "not_counting": "(không kể thời gian phát đề)",
}
_EN = {
    "q": "Question", "code": "Exam code", "part": "PART", "end": "— THE END —",
    "parts": {"mcq": "MULTIPLE CHOICE", "tf": "TRUE/FALSE", "short": "SHORT ANSWER"},
    "name": "Full name: ……………………………………………   Student ID: ……………………",
    "course": "Course", "course_code": "Course code", "subject": "Subject", "time": "Time allowed",
    "minutes": "minutes", "not_counting": "",
}
_DEFAULTS = {
    ("vju", False): {"org1": "ĐẠI HỌC QUỐC GIA HÀ NỘI", "org2": "TRƯỜNG ĐẠI HỌC VIỆT NHẬT",
                     "title": "ĐỀ THI KẾT THÚC HỌC PHẦN"},
    ("vju", True): {"org1": "VIETNAM NATIONAL UNIVERSITY, HANOI", "org2": "VNU VIETNAM JAPAN UNIVERSITY",
                    "title": "FINAL EXAMINATION"},
    ("school", False): {"org1": "SỞ GIÁO DỤC VÀ ĐÀO TẠO ……", "org2": "TRƯỜNG ……", "title": "ĐỀ KIỂM TRA"},
    ("school", True): {"org1": "DEPARTMENT OF EDUCATION AND TRAINING", "org2": "…… SCHOOL", "title": "TEST"},
}


def normalize(fmt: dict | None) -> dict:
    """A clean format dict (unknown keys dropped, defaults filled)."""
    f = fmt if isinstance(fmt, dict) else {}
    try:
        start = int(f.get("start") or 1)
    except (TypeError, ValueError):
        start = 1
    out = {
        "header": f.get("header") if f.get("header") in HEADERS else "none",
        "label": ":" if f.get("label") == ":" else ".",
        "start": min(max(start, 1), MAX_START),
        "english": bool(f.get("english")),
    }
    for k in FIELDS:
        out[k] = str(f.get(k) or "").strip()[:200]
    return out


def words(fmt: dict) -> dict:
    return _EN if fmt.get("english") else _VN


def number(fmt: dict, n: int, digits_only: bool = False) -> str:
    """The number of the n-th câu of a part (n from 1): "Câu 41: " / "41. "."""
    k = n + fmt["start"] - 1
    return f"{k}{fmt['label']} " if digits_only else f"{words(fmt)['q']} {k}{fmt['label']} "


def part_heading(fmt: dict, part_of: str, qt: str) -> str:
    w = words(fmt)
    return f"{w['part']} {part_of}. {w['parts'][qt]}"


# ── Đầu đề ──────────────────────────────────────────────────────────────────

def _no_borders(table) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "nil")
        borders.append(e)
    tbl_pr.append(borders)


def _line(p, text: str, bold: bool = False, italic: bool = False, size: float = 12, new_line: bool = True):
    from docx.shared import Pt

    if new_line and p.runs:
        p.add_run().add_break()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.size = Pt(size)
    r.font.name = "Times New Roman"
    return r


def _para(doc, align: str = "center"):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    p = doc.add_paragraph()
    p.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
    p.paragraph_format.space_after = Pt(4)
    return p


def header_elements(doc, fmt: dict, code: str) -> list:
    """Build the đầu đề at the end of `doc` and return its body elements
    (table + paragraphs) in order; the caller places them. fmt["header"]
    must not be "none"."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm

    w = words(fmt)
    d = {**_DEFAULTS[(fmt["header"], fmt["english"])], **{k: v for k, v in fmt.items() if k in FIELDS and v}}
    els = []

    table = doc.add_table(rows=1, cols=2)
    _no_borders(table)
    left, right = table.rows[0].cells
    left.width, right.width = Cm(8.5), Cm(8.5)
    lp, rp = left.paragraphs[0], right.paragraphs[0]
    for p in (lp, rp):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _line(lp, d.get("org1", ""))
    _line(lp, d.get("org2", ""), bold=True)
    if d.get("org3"):
        _line(lp, d["org3"], bold=True)
    _line(rp, d.get("title", "").upper(), bold=True)
    if d.get("term"):
        _line(rp, d["term"], bold=fmt["header"] == "school")
    if fmt["header"] == "school":
        if d.get("subject"):
            _line(rp, f"{w['subject']}: {d['subject']}", bold=True)
        if d.get("duration"):
            _line(rp, f"{w['time']}: {d['duration']} {w['minutes']}", italic=True)
            if w["not_counting"]:
                _line(rp, w["not_counting"], italic=True, size=11)
    els.append(table._tbl)

    if fmt["header"] == "vju":
        if d.get("subject") or d.get("subject_code"):
            p = _para(doc)
            text = f"{w['course']}: {d.get('subject', '')}"
            if d.get("subject_code"):
                text += f" – {w['course_code']}: {d['subject_code']}"
            _line(p, text, bold=True)
            els.append(p._p)
        if d.get("duration"):
            p = _para(doc)
            _line(p, f"{w['time']}: {d['duration']} {w['minutes']}", italic=True)
            els.append(p._p)

    p = _para(doc)
    _line(p, f"{w['code']}: {code}", bold=True, size=13)
    els.append(p._p)
    p = _para(doc, "left")
    _line(p, w["name"])
    els.append(p._p)
    return els
