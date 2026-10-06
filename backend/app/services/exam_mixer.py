"""
Trộn đề — build shuffled exam versions (mã đề) and print them to Word.

Rules (agreed 2026-09-29, for the "Mẫu 40 câu TN + Đúng/Sai" answer sheet):
  - every version has the SAME questions, only the order differs;
  - order is shuffled inside each part: trắc nghiệm (Phần I-II on the VJU
    sheet), then Đúng/Sai (Phần III), then trả lời ngắn (Phần IV);
  - trắc nghiệm options are shuffled unless the question says not to
    (shuffle_options=False) — options pinned with "#" stay in place;
  - Đúng/Sai statements a)–d) keep their order; trả lời ngắn has no options;
  - "for the sheet" (chấm bằng phiếu Mẫu 40): the sheet has room for at
    most 40 / 8 / 6 questions in the three parts and only 4 bubbles (A–D)
    per trắc nghiệm question — a question with more than 4 options is left out.
    Without it (đề chỉ để in, chấm tay / chấm khác) any number of questions
    and up to 8 options; such a bộ đề can't be attached to a kỳ thi.

YoungMix question groups (Word files, see question_io): a "<gN>" line starts
a group of consecutive questions —
    <g0> nothing shuffled inside   <g1> question order only
    <g2> options only              <g3> both
Groups (blocks) swap places with each other when questions are shuffled;
"<#gN>" pins a group to its position. Questions before any group line form
one block that behaves like <g3>. Both checkboxes (trộn câu / trộn đáp án)
still act as master switches.
"""
from __future__ import annotations

import io
import random
import re
from dataclasses import dataclass

from app.services.question_io import LETTERS, PART_NAME, PART_OF, QTYPES, ParsedQuestion

PART_LIMITS = {"mcq": 40, "tf": 8, "short": 6}
MAX_MCQ_OPTIONS = 4
MAX_VERSIONS = 24
MAX_FREE_PER_PART = 500     # đề không dùng phiếu — just a sanity cap


class MixError(ValueError):
    """Bad mixing request — the message is shown to the teacher as-is."""


@dataclass
class Snapshot:
    """A question as frozen into one version (options already shuffled)."""
    qtype:       str
    content:     str
    options:     list[dict]           # [{"text", "correct"?}]
    answer:      int = 0              # mcq: index of the correct option, -1 = not known (đề chỉ để in)
    answer_text: str | None = None    # short
    # Word source only (docx_mixer): the question's index in the parsed file and,
    # for trắc nghiệm, the original option shown at each position.
    src:         int | None = None
    perm:        list[int] | None = None

    def to_dict(self) -> dict:
        d = {"qtype": self.qtype, "content": self.content, "options": self.options,
             "answer": self.answer, "answer_text": self.answer_text}
        if self.src is not None:
            d["src"], d["perm"] = self.src, self.perm
        return d

    @staticmethod
    def from_dict(d: dict) -> "Snapshot":
        return Snapshot(d["qtype"], d["content"], d.get("options") or [], d.get("answer", 0), d.get("answer_text"),
                        d.get("src"), d.get("perm"))


def _rules(sheet) -> tuple[str, dict[str, int], int, int | None]:
    """(name, part limits, most options per trắc nghiệm, mã đề digit columns)
    of the answer sheet — `sheet` is a SheetLayout (sheet_layouts.py); None =
    the Mẫu 40 câu sheet, as before 2026-10-05."""
    if sheet is None:
        return "phiếu Mẫu 40 câu trắc nghiệm", PART_LIMITS, MAX_MCQ_OPTIONS, 4
    return f"phiếu \"{sheet.name}\"", sheet.limits, sheet.mcq_options, sheet.code_digits


def usable(q: ParsedQuestion, for_sheet: bool = True, sheet=None) -> bool:
    return not (for_sheet and q.qtype == "mcq" and len(q.options) > _rules(sheet)[2])


def check_counts(available: dict[str, int], counts: dict[str, int], for_sheet: bool = True, sheet=None) -> None:
    """Raise MixError if the requested numbers can't go on the sheet / aren't there."""
    if sum(counts.values()) == 0:
        raise MixError("Chọn ít nhất 1 câu hỏi")
    sheet_name, limits, _, _ = _rules(sheet)
    for qt in QTYPES:
        want = counts.get(qt, 0)
        name = f"Phần {PART_OF[qt]} ({PART_NAME[qt].lower()})"
        if want < 0:
            raise MixError(f"{name}: số câu không hợp lệ")
        if for_sheet and want > limits[qt]:
            raise MixError(f"{name}: {sheet_name} chỉ có {limits[qt]} câu nhưng đang chọn {want}. "
                           "Chọn \"Chỉ in đề\" nếu đề chỉ để in")
        if want > MAX_FREE_PER_PART:
            raise MixError(f"{name}: tối đa {MAX_FREE_PER_PART} câu nhưng đang chọn {want}")
        if want > available.get(qt, 0):
            raise MixError(f"{name}: chỉ có {available.get(qt, 0)} câu dùng được nhưng đang chọn {want}")


def sheet_problem(snaps: list["Snapshot"], code: str = "0", sheet=None) -> str | None:
    """Why this version can't be graded on the answer sheet (None = it can)."""
    _, limits, max_opts, _ = _rules(sheet)
    for qt in QTYPES:
        n = sum(1 for s in snaps if s.qtype == qt)
        if n > limits[qt]:
            return f"Phần {PART_OF[qt]} có {n} câu, phiếu chỉ có {limits[qt]}"
    if any(s.qtype == "mcq" and len(s.options) > max_opts for s in snaps):
        return f"có câu trắc nghiệm hơn {max_opts} đáp án, phiếu chỉ có {LETTERS[0]}–{LETTERS[max_opts - 1]}"
    if any(s.qtype == "mcq" and (s.answer is None or s.answer < 0) for s in snaps):
        return "có câu chưa có đáp án"
    if not code.isdigit():
        return f"mã đề \"{code}\" không phải số"
    digits = _rules(sheet)[3]
    if digits is not None and len(code) > digits:
        return f"mã đề {code} có {len(code)} chữ số, phiếu chỉ có {digits} ô mã đề"
    return None


_CODE_OK = re.compile(r"^[\w][\w .\-]*$")
MAX_CODE_LEN = 10


def version_codes(spec: str, n: int, for_sheet: bool = True, sheet=None) -> list[str]:
    """The mã đề of a new bộ đề, from what the teacher typed:
      - a list "A, B, C, D" / "Đề 1; Đề 2" → exactly those (n is ignored);
      - one code → n codes counting up from it: "101" → 101, 102, …
        (leading zeros kept: "001" → "002"), "MD01" → MD02, "A" → B, C, …
    On the "Mẫu 40 câu" sheet the mã đề is bubbled as digits, so for_sheet
    codes must be numbers of at most 3 digits (4 if the teacher starts with 4),
    and never more digits than the sheet's Mã đề field has columns. A sheet
    without a Mã đề field can't tell mã đề apart: 1 mã đề only."""
    spec = (spec or "").strip()
    if not spec:
        raise MixError("Nhập mã đề (vd 101, A, Đề 1, hoặc danh sách A, B, C, D)")
    if re.search(r"[,;]", spec):
        codes = [c.strip() for c in re.split(r"[,;]", spec) if c.strip()]
        if not 1 <= len(codes) <= MAX_VERSIONS:
            raise MixError(f"Số mã đề phải từ 1 đến {MAX_VERSIONS}")
    else:
        if not 1 <= n <= MAX_VERSIONS:
            raise MixError(f"Số mã đề phải từ 1 đến {MAX_VERSIONS}")
        if m := re.match(r"^(.*?)(\d+)$", spec):
            prefix, num = m.groups()
            codes = [prefix + str(int(num) + i).zfill(len(num)) for i in range(n)]
        elif m := re.match(r"^(.*?)([A-Za-z])$", spec):
            prefix, ch = m.groups()
            base = "A" if ch.isupper() else "a"
            if ord(ch) - ord(base) + n > 26:
                raise MixError(f"Không đủ chữ cái để tạo {n} mã đề từ \"{spec}\"")
            codes = [prefix + chr(ord(ch) + i) for i in range(n)]
        else:
            raise MixError("Mã đề phải kết thúc bằng số hoặc chữ cái để tự tăng (vd 101, A, Đề 1), "
                           "hoặc nhập danh sách cách nhau bằng dấu phẩy")
    for c in codes:
        if len(c) > MAX_CODE_LEN or not _CODE_OK.match(c):
            raise MixError(f"Mã đề \"{c}\" không hợp lệ (tối đa {MAX_CODE_LEN} ký tự, chỉ gồm chữ, số, dấu cách, - _ .)")
    if len({c.casefold() for c in codes}) != len(codes):
        raise MixError("Các mã đề bị trùng nhau")
    if for_sheet:
        sheet_name, _, _, digits = _rules(sheet)
        if digits is None:
            if len(codes) > 1:
                raise MixError(f"{sheet_name[0].upper()}{sheet_name[1:]} không có ô Mã đề nên chỉ trộn được 1 mã đề. "
                               "Chọn \"Chỉ in đề\" để in nhiều mã đề")
            return codes
        if not all(c.isdigit() for c in codes):
            raise MixError(f"Chấm bằng {sheet_name} thì mã đề phải là số (vd 101) vì phiếu chỉ tô được chữ số. "
                           "Chọn \"Chỉ in đề\" để dùng mã đề chữ")
        width = min(max(len(codes[0]), min(3, digits)), digits)
        if any(len(c) > width for c in codes):
            raise MixError(f"Mã đề vượt quá số chữ số trên phiếu (tối đa {digits} chữ số), hãy chọn mã đề nhỏ hơn")
    return codes


def _blocks(part: list[ParsedQuestion]) -> list[list[ParsedQuestion]]:
    """Consecutive questions of the same group form one block."""
    blocks: list[list[ParsedQuestion]] = []
    for q in part:
        if blocks and blocks[-1][0].group == q.group:
            blocks[-1].append(q)
        else:
            blocks.append([q])
    return blocks


def _order_part(part: list[ParsedQuestion], rng: random.Random, shuffle_questions: bool) -> list[ParsedQuestion]:
    blocks = _blocks(part)
    if shuffle_questions:
        # Swap block positions, except "<#gN>" blocks which stay where they are
        movable = [i for i, b in enumerate(blocks) if not b[0].group_fixed]
        order = movable[:]
        rng.shuffle(order)
        placed = blocks[:]
        for slot, src in zip(movable, order):
            placed[slot] = blocks[src]
        blocks = placed
        # Inside a block: shuffle unless the group says keep the order (<g0>, <g2>)
        blocks = [rng.sample(b, len(b)) if (b[0].group_mode is None or b[0].group_mode in (1, 3)) else b
                  for b in blocks]
    return [q for b in blocks for q in b]


def _shuffle_options(q: ParsedQuestion, rng: random.Random, enabled: bool) -> tuple[list[dict], int, list[int]]:
    """Shuffle a trắc nghiệm question's options; "fixed" ones keep their slot.
    Returns (options, index of the correct one or -1 if unknown, order)."""
    order = list(range(len(q.options)))
    if enabled and q.shuffle_options:
        movable = [i for i, o in enumerate(q.options) if not o.get("fixed")]
        shuffled = movable[:]
        rng.shuffle(shuffled)
        for slot, src in zip(movable, shuffled):
            order[slot] = src
    ans = order.index(q.answer) if q.answer is not None and q.answer >= 0 else -1
    return [{"text": q.options[i]["text"]} for i in order], ans, order


def build_versions(
    questions: list[ParsedQuestion],
    codes: list[str],
    *,
    shuffle_questions: bool = True,
    shuffle_options: bool = True,
    seed: int | None = None,
) -> dict[str, list[Snapshot]]:
    """{code: [Snapshot, …]} — trắc nghiệm, Đúng/Sai, trả lời ngắn in that order."""
    rng = random.Random(seed)
    out: dict[str, list[Snapshot]] = {}
    for code in codes:
        items: list[Snapshot] = []
        for qt in QTYPES:
            part = _order_part([q for q in questions if q.qtype == qt], rng, shuffle_questions)
            for q in part:
                if qt == "mcq":
                    opts, ans, order = _shuffle_options(q, rng, shuffle_options)
                    items.append(Snapshot("mcq", q.content, opts, ans, src=q.src_index,
                                          perm=order if q.src_index is not None else None))
                elif qt == "tf":
                    items.append(Snapshot("tf", q.content,
                                          [{"text": o["text"], "correct": bool(o.get("correct"))} for o in q.options],
                                          src=q.src_index))
                else:
                    items.append(Snapshot("short", q.content, [], 0, q.answer_text, src=q.src_index))
        out[code] = items
    return out


def answer_key(snaps: list[Snapshot]) -> dict[str, list]:
    """Human-readable key of one version, per part:
    {"mcq": ["A", …], "tf": [["Đ","S","Đ","S"], …], "short": ["-1,5", …]}"""
    key: dict[str, list] = {"mcq": [], "tf": [], "short": []}
    for s in snaps:
        if s.qtype == "mcq":
            key["mcq"].append(LETTERS[s.answer] if s.answer is not None and s.answer >= 0 else "?")
        elif s.qtype == "tf":
            key["tf"].append(["Đ" if o.get("correct") else "S" for o in s.options])
        else:
            key["short"].append(s.answer_text or "")
    return key


# ── Answer key in the grading labels of the "Mẫu 40 câu" sheet ──────────────
# Phần I-II (trắc nghiệm) → trc_nghim_abcd1..40        ("A".."D")
# Phần III (Đúng/Sai)     → ng_sai_cu1..32, 4 per câu  ("Đ" / "S")
# Phần IV (trả lời ngắn)  → the 6 signed-decimal fields, in sheet order; the OMR reads them
#            with a "." decimal point ("-1.5"), so the key uses "." too.

def grading_key(snaps: list[Snapshot], sheet) -> dict[str, str]:
    """The answer key of a version in the grading labels of its answer sheet
    (SheetLayout: trắc nghiệm, Đúng/Sai 4 per câu, trả lời ngắn field keys)."""
    key = answer_key(snaps)
    out: dict[str, str] = {}
    for label, letter in zip(sheet.mcq, key["mcq"]):
        if letter != "?":             # not known → not graded
            out[label] = letter
    for i, statements in enumerate(key["tf"]):
        for j, v in enumerate(statements):
            if i * 4 + j < len(sheet.tf):
                out[sheet.tf[i * 4 + j]] = v
    for label, value in zip(sheet.short, key["short"]):
        out[label] = value.replace(",", ".")
    return out


def grading_key_mau40(snaps: list[Snapshot], short_labels: list[str]) -> dict[str, str]:
    key = answer_key(snaps)
    out: dict[str, str] = {}
    for i, letter in enumerate(key["mcq"], start=1):
        if letter != "?":             # not known → not graded
            out[f"trc_nghim_abcd{i}"] = letter
    for i, statements in enumerate(key["tf"]):
        for j, v in enumerate(statements, start=1):
            out[f"ng_sai_cu{i * 4 + j}"] = v
    for label, value in zip(short_labels, key["short"]):
        out[label] = value.replace(",", ".")
    return out


# ── Word output ──────────────────────────────────────────────────────────────

def _new_doc():
    from docx import Document
    from docx.shared import Cm, Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(1.8)
        s.left_margin = s.right_margin = Cm(2)
    return doc


def _center(doc, text: str, bold: bool = False, size: float | None = None):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = bold
    if size:
        r.font.size = Pt(size)
    return p


def version_docx(title: str, code: str, snaps: list[Snapshot], subtitle: str = "",
                 assets: dict | None = None) -> bytes:
    """The exam students receive — no answers marked. Formulas/pictures
    ("[[ct:<id>]]") are written as the real objects from `assets`."""
    from app.services.rich_objects import add_rich

    doc = _new_doc()
    _center(doc, title.upper(), bold=True, size=14)
    if subtitle:
        _center(doc, subtitle)
    _center(doc, f"Mã đề: {code}", bold=True, size=13)
    doc.add_paragraph("Họ và tên: ……………………………………………   Số báo danh: ……………………")

    for qt in QTYPES:
        part = [s for s in snaps if s.qtype == qt]
        if not part:
            continue
        doc.add_paragraph().add_run(f"PHẦN {PART_OF[qt]}. {PART_NAME[qt]}").bold = True
        for n, s in enumerate(part, start=1):
            p = doc.add_paragraph()
            p.add_run(f"Câu {n}. ").bold = True
            add_rich(doc, p, s.content, assets)
            if qt == "mcq":
                for i, o in enumerate(s.options):
                    op = doc.add_paragraph()
                    op.paragraph_format.left_indent = None
                    op.add_run(f"{LETTERS[i]}. ").bold = True
                    add_rich(doc, op, o["text"], assets)
            elif qt == "tf":
                for i, o in enumerate(s.options):
                    op = doc.add_paragraph()
                    op.add_run(f"{LETTERS[i].lower()}) ")
                    add_rich(doc, op, o["text"], assets)
    _center(doc, "— HẾT —", bold=True)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def answer_key_docx(title: str, versions: dict[str, list[Snapshot]]) -> bytes:
    """Đáp án of every mã đề: one table per part, one row per mã đề."""
    doc = _new_doc()
    _center(doc, f"ĐÁP ÁN: {title.upper()}", bold=True, size=14)
    keys = {code: answer_key(snaps) for code, snaps in versions.items()}
    codes = list(keys)

    def table(header: list[str], rows: list[list[str]]):
        t = doc.add_table(rows=1 + len(rows), cols=len(header))
        t.style = "Table Grid"
        for i, h in enumerate(header):
            t.rows[0].cells[i].text = h
            t.rows[0].cells[i].paragraphs[0].runs[0].bold = True
        for r, row in enumerate(rows, start=1):
            for i, v in enumerate(row):
                t.rows[r].cells[i].text = v

    for qt in QTYPES:
        n = len(keys[codes[0]][qt]) if codes else 0
        if not n:
            continue
        doc.add_paragraph().add_run(f"PHẦN {PART_OF[qt]}. {PART_NAME[qt]}").bold = True
        # trắc nghiệm can be 40 wide — split into chunks of 10 questions per table
        step = 10 if qt == "mcq" else n
        for start in range(0, n, step):
            idx = range(start, min(start + step, n))
            header = ["Mã đề"] + [f"Câu {i + 1}" for i in idx]
            rows = []
            for code in codes:
                vals = keys[code][qt]
                rows.append([code] + ["".join(vals[i]) if qt == "tf" else vals[i] for i in idx])
            table(header, rows)
            doc.add_paragraph()

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
