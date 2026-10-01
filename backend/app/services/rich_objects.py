"""
Công thức và hình ảnh trong câu hỏi (2026-10-01).

"t muốn nó phải hiện đúng đủ và hiện được": a Word đề's MathType formulas
(w:object, Equation.DSMT4), pictures (w:drawing / w:pict) and Word equations
(m:oMath) used to be dropped on import, leaving options like "B." empty.
Now each one becomes a token in the question text,

    "Căn bậc hai số học của [[ct:9f2c…]] là:"

and the object itself is an asset kept by id (question_assets table):

  collect(el, part)   → the asset of a Word element: its XML and the package
                        parts it points at (MathType .bin + preview, image…).
                        The id is a hash of that content, so the same formula
                        is one asset and re-importing a file finds the same ids.
  embed(doc, par, a)  → put the asset back into another Word file, the
                        ORIGINAL object (a MathType formula stays editable).
  render_svgs(assets) → an SVG of each for the web: every asset on its own
                        page of one Word file, LibreOffice turns it into a PDF
                        in a single run, PyMuPDF crops each page to the ink and
                        writes SVG with text as paths (browsers can't draw WMF).

Without LibreOffice (render_svgs returns {}) nothing is lost: Word exports
and mixed đề still carry the real objects; only the web shows a placeholder.
"""
from __future__ import annotations

import base64
import copy
import hashlib
import io
import logging
import os
import re
import shutil
import subprocess
import tempfile
import weakref

logger = logging.getLogger(__name__)

TOKEN_RE = re.compile(r"\[\[ct:([0-9a-f]{8,40})\]\]")

NS = {
    "w":  "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "m":  "http://schemas.openxmlformats.org/officeDocument/2006/math",
    "v":  "urn:schemas-microsoft-com:vml",
    "o":  "urn:schemas-microsoft-com:office:office",
    "r":  "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "a":  "http://schemas.openxmlformats.org/drawingml/2006/main",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
}
R_NS = NS["r"]

# The VML "picture frame" shape type MathType objects refer to (#_x0000_t75).
# Word writes it once per document; an object copied alone needs its own.
_SHAPETYPE_75 = (
    '<v:shapetype xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office" '
    'id="_x0000_t75" coordsize="21600,21600" o:spt="75" o:preferrelative="t" path="m@4@5l@4@11@9@11@9@5xe" '
    'filled="f" stroked="f"><v:stroke joinstyle="miter"/><v:formulas>'
    '<v:f eqn="if lineDrawn pixelLineWidth 0"/><v:f eqn="sum @0 1 0"/><v:f eqn="sum 0 0 @1"/>'
    '<v:f eqn="prod @2 1 2"/><v:f eqn="prod @3 21600 pixelWidth"/><v:f eqn="prod @3 21600 pixelHeight"/>'
    '<v:f eqn="sum @0 0 1"/><v:f eqn="prod @6 1 2"/><v:f eqn="prod @7 21600 pixelWidth"/>'
    '<v:f eqn="sum @8 21600 0"/><v:f eqn="prod @7 21600 pixelHeight"/><v:f eqn="sum @10 21600 0"/>'
    '</v:formulas><v:path o:extrusionok="f" gradientshapeok="t" o:connecttype="rect"/>'
    '<o:lock v:ext="edit" aspectratio="t"/></v:shapetype>'
)


def token(asset_id: str) -> str:
    return f"[[ct:{asset_id}]]"


def ids_in(text: str) -> list[str]:
    return TOKEN_RE.findall(text or "")


def plain(text: str, placeholder: str = "[công thức]") -> str:
    """Text without tokens (Moodle .txt export, short labels)."""
    return TOKEN_RE.sub(placeholder, text or "")


def _q(tag: str) -> str:
    pre, name = tag.split(":")
    return f"{{{NS[pre]}}}{name}"


def _size_pt(el) -> tuple[float, float]:
    """Width/height of a Word object in points (VML style or DrawingML extent)."""
    sh = el.find(".//" + _q("v:shape"))
    if sh is not None:
        style = sh.get("style", "")

        def dim(key: str) -> float:
            m = re.search(key + r"\s*:\s*([\d.]+)\s*(pt|in|px|cm|mm)?", style)
            if not m:
                return 0.0
            v, unit = float(m.group(1)), (m.group(2) or "pt")
            return v * {"pt": 1, "in": 72, "px": 0.75, "cm": 28.3465, "mm": 2.83465}[unit]

        return dim("width"), dim("height")
    ext = el.find(".//" + _q("wp:extent"))
    if ext is not None:
        return int(ext.get("cx", 0)) / 12700, int(ext.get("cy", 0)) / 12700
    return 0.0, 0.0


def collect(el, part) -> dict | None:
    """The asset of a w:object / w:drawing / w:pict / m:oMath element of a
    Word part (python-docx Part, to resolve r:id). None if it's empty."""
    from lxml import etree

    kind = {_q("w:object"): "object", _q("w:drawing"): "drawing", _q("w:pict"): "pict",
            _q("m:oMath"): "omath"}.get(el.tag)
    if kind is None:
        return None
    el = copy.deepcopy(el)
    parts: dict[str, dict] = {}
    for node in el.iter():
        for attr, rid in list(node.attrib.items()):
            if not attr.startswith(f"{{{R_NS}}}") or rid in parts:
                continue
            rel = part.rels.get(rid)
            if rel is None or rel.is_external:
                continue
            p = rel.target_part
            parts[rid] = {"blob": p.blob, "content_type": p.content_type, "reltype": rel.reltype,
                          "ext": str(p.partname).rsplit(".", 1)[-1].lower()}
    if kind != "omath" and not parts:
        return None             # a shape with nothing to show
    # one picture-frame definition per object, so it renders copied alone
    if kind in ("object", "pict") and 'type="#_x0000_t75"' in etree.tostring(el).decode() \
            and el.find(".//" + _q("v:shapetype")) is None:
        sh = el.find(".//" + _q("v:shape"))
        if sh is not None:
            sh.addprevious(etree.fromstring(_SHAPETYPE_75))
    w, h = _size_pt(el)
    xml = etree.tostring(el, encoding="unicode")
    digest = hashlib.sha1()
    digest.update(kind.encode())
    if parts:
        for rid in sorted(parts):
            digest.update(parts[rid]["blob"])
        digest.update(f"{w:.1f}x{h:.1f}".encode())
    else:
        digest.update(xml.encode())
    return {"id": digest.hexdigest()[:24], "kind": kind, "xml": xml, "parts": parts, "w_pt": w, "h_pt": h}


# ── Back into a Word file ────────────────────────────────────────────────────

# per target document: parts already added (by blob) and the last shape number
_EMBED_STATE: weakref.WeakKeyDictionary = weakref.WeakKeyDictionary()


def _partname_template(ext: str, reltype: str) -> str:
    if reltype.endswith("/oleObject") or ext == "bin":
        return "/word/embeddings/oleObject%d." + ext
    return "/word/media/image%d." + ext


def embed(doc, paragraph, asset: dict) -> None:
    """Append the asset's object to `paragraph` of python-docx `doc`."""
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    from lxml import etree

    el = etree.fromstring(asset["xml"])
    state = _EMBED_STATE.setdefault(doc.part, {"parts": {}, "shape_no": [3000]})
    cache, counter = state["parts"], state["shape_no"]          # same blob → same part
    new_ids: dict[str, str] = {}
    for rid, p in asset["parts"].items():
        key = (hashlib.sha1(p["blob"]).hexdigest(), p["reltype"])
        if key not in cache:
            name = doc.part.package.next_partname(_partname_template(p["ext"], p["reltype"]))
            cache[key] = Part(PackURI(name), p["content_type"], p["blob"], doc.part.package)
        new_ids[rid] = doc.part.relate_to(cache[key], p["reltype"])
    for node in el.iter():
        for attr, val in list(node.attrib.items()):
            if attr.startswith(f"{{{R_NS}}}") and val in new_ids:
                node.set(attr, new_ids[val])
    # shape ids must be unique in a document
    counter[0] += 1
    sid = f"_x0000_i{counter[0]}"
    for sh in el.iter(_q("v:shape")):
        sh.set("id", sid)
    for ole in el.iter(_q("o:OLEObject")):
        ole.set("ShapeID", sid)
    for pr in el.iter(_q("wp:docPr")):
        pr.set("id", str(counter[0]))
    if asset["kind"] == "omath":
        paragraph._p.append(el)
    else:
        paragraph.add_run()._r.append(el)


def add_rich(doc, paragraph, text: str, assets: dict[str, dict] | None, bold: bool = False):
    """Write `text` into `paragraph`, tokens as their real objects (or
    "[công thức]" when the asset isn't known)."""
    pos = 0
    for m in TOKEN_RE.finditer(text or ""):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()]).bold = bold
        a = (assets or {}).get(m.group(1))
        if a is not None:
            try:
                embed(doc, paragraph, a)
            except Exception:
                logger.exception("embedding asset %s failed", m.group(1))
                paragraph.add_run("[công thức]")
        else:
            paragraph.add_run("[công thức]")
        pos = m.end()
    if pos < len(text or ""):
        paragraph.add_run(text[pos:]).bold = bold


# ── SVG for the web ──────────────────────────────────────────────────────────

def soffice_path() -> str | None:
    for cand in (os.environ.get("SOFFICE_PATH"), shutil.which("soffice"), shutil.which("libreoffice"),
                 "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if cand and os.path.exists(cand):
            return cand
    return None


def render_svgs(assets: list[dict], timeout: int = 300) -> dict[str, str]:
    """{asset id: svg} for the ones LibreOffice could draw ({} without it)."""
    if not assets:
        return {}
    soffice = soffice_path()
    if soffice is None:
        return {}
    try:
        import pymupdf
    except ImportError:
        try:
            import fitz as pymupdf      # older PyMuPDF
        except ImportError:
            logger.warning("PyMuPDF is not installed: formulas can't be shown on the web")
            return {}
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document()
    body = doc.element.body
    for ch in list(body):
        body.remove(ch)

    def sect(w_pt: float, h_pt: float):
        s = OxmlElement("w:sectPr")
        pg = OxmlElement("w:pgSz")
        pg.set(qn("w:w"), str(int(max(w_pt, 20) * 20 + 200)))
        pg.set(qn("w:h"), str(int(max(h_pt, 10) * 20 + 200)))
        s.append(pg)
        mar = OxmlElement("w:pgMar")
        for a in ("top", "bottom", "left", "right"):
            mar.set(qn(f"w:{a}"), "60")
        for a in ("header", "footer", "gutter"):
            mar.set(qn(f"w:{a}"), "0")
        s.append(mar)
        return s

    from docx.text.paragraph import Paragraph
    for k, a in enumerate(assets):
        w, h = (a["w_pt"], a["h_pt"]) if a["w_pt"] and a["h_pt"] else (430, 120)
        p = OxmlElement("w:p")
        body.append(p)
        par = Paragraph(p, doc._body)
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.space_before = Pt(0)
        embed(doc, par, a)
        if k < len(assets) - 1:
            p.get_or_add_pPr().append(sect(w, h))
    body.append(sect(*((assets[-1]["w_pt"], assets[-1]["h_pt"]) if assets[-1]["w_pt"] else (430, 120))))

    out: dict[str, str] = {}
    with tempfile.TemporaryDirectory(prefix="vju-ct-") as tmp:
        src = os.path.join(tmp, "objects.docx")
        doc.save(src)
        profile = "file://" + os.path.join(tempfile.gettempdir(), "vju-soffice-profile")
        try:
            subprocess.run([soffice, "--headless", "--norestore", f"-env:UserInstallation={profile}",
                            "--convert-to", "pdf", "--outdir", tmp, src],
                           stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=timeout, check=False)
        except (OSError, subprocess.TimeoutExpired):
            logger.exception("LibreOffice failed to draw %d formulas", len(assets))
            return {}
        pdf_path = os.path.join(tmp, "objects.pdf")
        if not os.path.exists(pdf_path):
            logger.warning("LibreOffice produced no PDF for %d formulas", len(assets))
            return {}
        pdf = pymupdf.open(pdf_path)
        if pdf.page_count != len(assets):
            logger.warning("formula pages: expected %d, got %d", len(assets), pdf.page_count)
        for a, page in zip(assets, pdf):
            box = None
            for _kind, r in page.get_bboxlog():
                rect = pymupdf.Rect(r)
                if rect.is_empty or rect.width > page.rect.width - 1 and rect.height > page.rect.height - 1:
                    continue        # the page background
                box = rect if box is None else box | rect
            if box is None:
                continue
            box = (box + (-0.6, -0.6, 0.6, 0.6)) & page.rect
            page.set_cropbox(box)
            svg = page.get_svg_image(text_as_path=True)
            # size in points, so a formula is as big next to the text as in Word
            svg = re.sub(r'(<svg[^>]*?)\swidth="([\d.]+)"\s+height="([\d.]+)"',
                         lambda m: f'{m.group(1)} width="{m.group(2)}pt" height="{m.group(3)}pt"', svg, count=1)
            out[a["id"]] = svg
    return out


def _wrap_png(png: bytes, w_pt: float, h_pt: float) -> str:
    b64 = base64.b64encode(png).decode()
    return (f'<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" '
            f'width="{w_pt:.1f}pt" height="{h_pt:.1f}pt" viewBox="0 0 {w_pt:.1f} {h_pt:.1f}">'
            f'<image width="{w_pt:.1f}" height="{h_pt:.1f}" preserveAspectRatio="xMidYMid meet" '
            f'xlink:href="data:image/png;base64,{b64}"/></svg>')


def _picture_blob(asset: dict) -> tuple[bytes, str] | None:
    """(blob, ext) of a picture asset's image, if it has exactly one."""
    imgs = [p for p in asset["parts"].values() if p["reltype"].endswith("/image")]
    return (imgs[0]["blob"], imgs[0]["ext"]) if len(imgs) == 1 else None


def picture_svgs(assets: list[dict]) -> dict[str, str]:
    """Pictures the browser can show as they are (png/jpg/gif), no LibreOffice."""
    from PIL import Image

    out = {}
    for a in assets:
        pic = _picture_blob(a) if a["kind"] in ("drawing", "pict") else None
        if pic is None or pic[1] not in ("png", "jpg", "jpeg", "gif", "bmp"):
            continue
        blob, _ = pic
        try:
            im = Image.open(io.BytesIO(blob))
            buf = io.BytesIO()
            im.save(buf, "PNG")
        except Exception:
            continue
        w, h = (a["w_pt"], a["h_pt"]) if a["w_pt"] else (im.width * 0.75, im.height * 0.75)
        out[a["id"]] = _wrap_png(buf.getvalue(), w, h)
    return out


def vector_picture_svgs(assets: list[dict], timeout: int = 300) -> dict[str, str]:
    """EMF/WMF pictures LibreOffice draws blank inside a page: converted on
    their own to a large PNG, trimmed to the ink."""
    from PIL import Image, ImageOps

    soffice = soffice_path()
    todo = [(a, _picture_blob(a)) for a in assets if a["kind"] in ("drawing", "pict")]
    todo = [(a, pic) for a, pic in todo if pic and pic[1] in ("emf", "wmf")]
    if soffice is None or not todo:
        return {}
    out: dict[str, str] = {}
    with tempfile.TemporaryDirectory(prefix="vju-pic-") as tmp:
        files = []
        for k, (a, (blob, ext)) in enumerate(todo):
            path = os.path.join(tmp, f"p{k}.{ext}")
            open(path, "wb").write(blob)
            files.append(path)
        profile = "file://" + os.path.join(tempfile.gettempdir(), "vju-soffice-profile")
        flt = 'png:draw_png_Export:{"PixelWidth":{"type":"long","value":"3000"},"PixelHeight":{"type":"long","value":"4243"}}'
        try:
            subprocess.run([soffice, "--headless", "--norestore", f"-env:UserInstallation={profile}",
                            "--convert-to", flt, "--outdir", tmp, *files],
                           stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=timeout, check=False)
        except (OSError, subprocess.TimeoutExpired):
            return {}
        for k, (a, _) in enumerate(todo):
            png = os.path.join(tmp, f"p{k}.png")
            if not os.path.exists(png):
                continue
            im = Image.open(png).convert("RGB")
            box = ImageOps.invert(im.convert("L")).getbbox()
            if not box:
                continue
            buf = io.BytesIO()
            im.crop(box).save(buf, "PNG", optimize=True)
            out[a["id"]] = _wrap_png(buf.getvalue(), a["w_pt"] or im.width * 0.75, a["h_pt"] or im.height * 0.75)
    return out


def svgs_for(assets: list[dict]) -> dict[str, str]:
    """Best web picture of each asset: plain images as they are, the rest via
    LibreOffice (one run), and a second try for pictures drawn blank."""
    out = picture_svgs(assets)
    rest = [a for a in assets if a["id"] not in out]
    out.update(render_svgs(rest))
    missing = [a for a in rest if a["id"] not in out]
    if missing:
        out.update(vector_picture_svgs(missing))
    return out


# ── Serialising for the database ─────────────────────────────────────────────

def parts_to_json(parts: dict[str, dict]) -> dict:
    return {rid: {**{k: v for k, v in p.items() if k != "blob"}, "blob": base64.b64encode(p["blob"]).decode()}
            for rid, p in parts.items()}


def parts_from_json(data: dict) -> dict[str, dict]:
    return {rid: {**{k: v for k, v in p.items() if k != "blob"}, "blob": base64.b64decode(p["blob"])}
            for rid, p in data.items()}
