"""
Import / export for the question bank (Ngân hàng câu hỏi).

Two file formats:

* Aiken .txt — what Moodle exports (sample: questions-AET2015…txt):

      Question text (may span several lines)
      A) option
      B) option
      ANSWER: A

  Aiken only has multiple-choice questions, so it only carries Phần I-II.

* Word .docx — YoungMix-style, so teachers can reuse the Word files they
  already prepare for youngmix.vn, extended with the three kinds of
  question on the VJU answer sheet ("Mẫu 40"), named as ON THE SHEET:
    - "PHẦN I-II" / "PHẦN III" / "PHẦN IV" headings switch the question kind
      (no heading = trắc nghiệm). A heading naming the kind ("… ĐÚNG/SAI",
      "… TRẢ LỜI NGẮN", "… TRẮC NGHIỆM") wins over its number, so files
      numbered the Bộ GD way (Phần I TN / II Đúng-Sai / III trả lời ngắn)
      read right too; bare numbers are then settled by the questions
      themselves (a)–d) statements → Đúng/Sai, a numeric "Đáp án:" with
      no options → trả lời ngắn), see _part_qtype / flush.
    - a question starts with "Câu 1." / "Question 1:" / "1." …
    - trắc nghiệm (Phần I-II): options start with "A." "B." … (several may
      share one line, separated by tabs); "#A." pins that option in place
      when shuffling; the correct option's letter is red or underlined
      (or a line "Đáp án: B")
    - Đúng/Sai (Phần III): 4 statements "a)" "b)" "c)" "d)"; the TRUE ones
      are red or underlined (or a line "Đáp án: Đ S Đ S")
    - trả lời ngắn (Phần IV): a line "Đáp án: -1,5" (max 4 characters, as
      bubbled on the sheet)
    - an optional "<g0>".."<g3>" line before a group controls shuffling:
      g0/g1 = keep option order, g2/g3 = options may be shuffled

Both parsers return (questions, warnings) and never raise on bad input —
unparseable items are skipped with a warning naming the question, so the
UI can show a preview before anything is saved.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

LETTERS = "ABCDEFGH"


QTYPES = ("mcq", "tf", "short")
PART_OF   = {"mcq": "I-II", "tf": "III", "short": "IV"}   # as printed on the VJU sheet
PART_NAME = {"mcq": "TRẮC NGHIỆM", "tf": "ĐÚNG/SAI", "short": "TRẢ LỜI NGẮN"}


@dataclass
class ParsedQuestion:
    content: str
    options: list[dict] = field(default_factory=list)   # [{"text", "fixed", ("correct" for tf)}]
    answer:  int = 0                                    # mcq: index of the correct option
    shuffle_options: bool = True
    qtype:   str = "mcq"                                # "mcq" | "tf" | "short"
    answer_text: str | None = None                      # short: e.g. "-1,5"
    # YoungMix question groups (Word files only; not kept in the question bank):
    # the "<gN>" / "<#gN>" line the question sits under. None = no group line.
    group:       int | None = None                      # sequential id of the group in the file
    group_mode:  int | None = None                      # 0 none, 1 question order, 2 options, 3 both
    group_fixed: bool = False                           # "<#gN>": the group keeps its position
    where:       str = field(default="", compare=False) # where it is in the file, for messages ("Câu 12, dòng 245")
    # Word files only: position in parse_docx's result and where the question
    # sits in the document (body element indexes, option offsets), so the
    # mixer can rearrange the ORIGINAL paragraphs — formulas, pictures and
    # formatting included (see docx_mixer.py). Not kept in the question bank.
    src_index:   int | None = field(default=None, compare=False)
    src:         dict | None = field(default=None, compare=False, repr=False)
    # Word files: the formulas/pictures its text refers to as "[[ct:<id>]]"
    # ({id: asset}, see rich_objects.py); stored as question_assets rows.
    assets:      dict = field(default_factory=dict, compare=False, repr=False)


_SHORT_ANSWER = re.compile(r"^-?\d+(,\d+)?$")


def normalize_short_answer(raw: str) -> str | None:
    """A trả lời ngắn (Phần IV) answer as it can be bubbled on the sheet: optional "-",
    digits, optional decimal comma — at most 4 characters ("-1,5", "2025").
    "." is accepted and turned into ",". Returns None when it can't be
    bubbled."""
    s = "".join(str(raw).split()).replace(".", ",")
    if not s or len(s) > 4 or not _SHORT_ANSWER.match(s):
        return None
    return s


def _is_true_false(options: list[dict]) -> bool:
    texts = [o["text"].strip().lower() for o in options]
    return texts in (["true", "false"], ["đúng", "sai"], ["false", "true"], ["sai", "đúng"])


def _short(text: str, n: int = 60) -> str:
    text = " ".join(text.split())
    return text if len(text) <= n else text[: n - 1] + "…"


# ── Aiken (.txt) ─────────────────────────────────────────────────────────────

_AIKEN_OPTION = re.compile(r"^([A-H])[\).]\s*(.*)$")
_AIKEN_ANSWER = re.compile(r"^ANSWER\s*:\s*(.*)$", re.IGNORECASE)


def parse_aiken(text: str) -> tuple[list[ParsedQuestion], list[str]]:
    questions: list[ParsedQuestion] = []
    warnings:  list[str] = []

    content_lines: list[str] = []
    options: list[dict] = []
    n = 0        # 1-based position of the item in the file, for warnings
    start = 0    # line the current item starts on

    def reset() -> None:
        content_lines.clear()
        options.clear()

    for lineno, raw in enumerate(text.lstrip("\ufeff").splitlines(), 1):
        line = raw.strip()
        if not line:
            continue

        m_ans = _AIKEN_ANSWER.match(line)
        if m_ans:
            n += 1
            where   = f"Câu {n} (dòng {start or lineno})"
            content = "\n".join(content_lines).strip()
            letter  = m_ans.group(1).strip().upper()[:1]
            if not content:
                warnings.append(f"{where}: thiếu nội dung câu hỏi, bỏ qua")
            elif len(options) < 2:
                warnings.append(f"{where} \"{_short(content)}\": không có phương án A/B/C/D"
                                f"{' và đáp án trống' if not letter else ''}, không phải câu trắc nghiệm "
                                "(trên Moodle có thể là câu tự luận, điền từ hoặc nối cặp) nên không trộn được, bỏ qua")
            elif not letter or LETTERS.find(letter) < 0 or LETTERS.index(letter) >= len(options):
                warnings.append(f"{where} \"{_short(content)}\": dòng ANSWER không hợp lệ, bỏ qua")
            else:
                questions.append(ParsedQuestion(
                    content=content,
                    options=[dict(o) for o in options],
                    answer=LETTERS.index(letter),
                    shuffle_options=not _is_true_false(options),
                    where=where,
                ))
            reset()
            start = 0
            continue

        m_opt = _AIKEN_OPTION.match(line)
        # An option must continue the A, B, C… sequence; otherwise the line
        # is part of the question text (e.g. a question starting with "A.")
        if m_opt and content_lines and LETTERS.index(m_opt.group(1)) == len(options):
            options.append({"text": m_opt.group(2).strip(), "fixed": False})
            continue

        if options:
            # Text after options but before ANSWER: continuation of last option
            options[-1]["text"] = (options[-1]["text"] + " " + line).strip()
        else:
            if not content_lines:
                start = lineno
            content_lines.append(line)

    if content_lines or options:
        warnings.append(f"Cuối file (dòng {start}): câu \"{_short(' '.join(content_lines))}\" thiếu dòng ANSWER, bỏ qua")

    return questions, warnings


def apply_picked_answers(questions: list[ParsedQuestion], answers_json: str | None) -> None:
    """Đáp án the teacher picked on the page for trắc nghiệm with none marked
    in the file: answers_json = {"<index in questions>": option}. Bad keys or
    out-of-range options are ignored; malformed JSON raises ValueError."""
    if not answers_json:
        return
    import json
    for key, ans in (json.loads(answers_json) or {}).items():
        try:
            q = questions[int(key)]
        except (ValueError, IndexError):
            continue
        if q.qtype == "mcq" and isinstance(ans, int) and 0 <= ans < len(q.options):
            q.answer = ans


def mcq_rows(questions: list[ParsedQuestion]) -> list[dict]:
    """Every trắc nghiệm of a file, for the page's "chọn đáp án" list and
    its file đáp án. i = index in `questions` (what answers_json is keyed
    by); mcq_no = its place among the file's trắc nghiệm (1 = first), the
    "#" the teacher sees, and what an answer file or a VJU answer key
    ("Phần I-II, câu 5") is matched by; answer = -1 when the file marks none."""
    rows = []
    for i, q in enumerate(questions):
        if q.qtype != "mcq":
            continue
        rows.append({"i": i, "mcq_no": len(rows) + 1, "where": q.where,
                     "number": (q.src or {}).get("number", "").strip(),
                     "content": q.content[:300], "options": [o["text"][:120] for o in q.options],
                     "answer": q.answer})
    return rows


def export_aiken(questions: list[ParsedQuestion]) -> str:
    """Aiken has no notion of pinned options — `fixed` is dropped — and only
    multiple-choice questions: Đúng/Sai and trả lời ngắn are left out, and so
    are questions with no answer yet (Aiken needs an ANSWER line)."""
    blocks = []
    for q in questions:
        if q.qtype != "mcq" or q.answer is None or q.answer < 0:
            continue
        # Moodle .txt can't hold a formula/picture: it says where one was
        lines = [TOKEN_IDS.sub("[công thức]", q.content).strip()]
        lines += [f"{LETTERS[i]}) {TOKEN_IDS.sub('[công thức]', o['text'])}" for i, o in enumerate(q.options)]
        lines.append(f"ANSWER: {LETTERS[q.answer]}")
        blocks.append("\n".join(lines))
    return "\n\n".join(blocks) + "\n"




# ── Word (.docx), YoungMix-style + the 3 parts of the answer sheet ──────────

# "Câu 1." / "Question 1:" … or just "1." / "1)" (many teachers' files number that way)
_DOCX_QUESTION  = re.compile(r"^\s*(?:(?:Câu|Cau|Question)\s*\d+\s*[:.)\-]?|\d{1,3}\s*[.)](?=\s|\S))\s*", re.IGNORECASE)
_DOCX_GROUP     = re.compile(r"^\s*<\s*(#?)\s*g([0-3])\s*>\s*$", re.IGNORECASE)
_DOCX_PART      = re.compile(r"^\s*ph[ầa]n\s+(i\s*[-–]\s*ii|iv|iii|ii|i|4|3|2|1)\b(.*)$", re.IGNORECASE)
_DOCX_ANSWER    = re.compile(r"^\s*(?:đáp\s*án|đáp\s*số|answer)\s*[:：]\s*(.*)$", re.IGNORECASE)
# Marker at line start, or after a tab / 2+ spaces (several on one line)
_DOCX_OPTION    = re.compile(r"(?:^|\t|\s{2,})\s*(#?)([A-H])[.)]\s")    # trắc nghiệm: A. B. …
_DOCX_STATEMENT = re.compile(r"(?:^|\t|\s{2,})\s*(#?)([a-h])[.)]\s")    # Đúng/Sai:   a) b) …
# Section headings between groups of questions ("I/ CĂN THỨC", "B. PHẦN HÌNH HỌC"):
# not part of any question — they used to be glued onto the last option.
_DOCX_SECTION   = re.compile(r"^\s*(?:[IVX]{1,5}\s*[/.)]\s*\S|[A-Z]\s*[.)/]\s*PH[ẦA]N\b)")
# End of the đề: "HẾT", "--- HẾT ---", or a long line of dashes
_DOCX_END       = re.compile(r"^[\s\-–—_=*.]*(?:HẾT|HET)[\s\-–—_=*.]*$|^\s*[-–—_=*]{5,}\s*$", re.IGNORECASE)
# Bare numbers, VJU sheet numbering (Phần I-II TN, III Đúng/Sai, IV trả lời ngắn)
_PART_QTYPE = {"i": "mcq", "1": "mcq", "i-ii": "mcq", "ii": "mcq", "2": "mcq",
               "iii": "tf", "3": "tf", "iv": "short", "4": "short"}


def _part_qtype(number: str, rest: str) -> str:
    """Kind of question under a "PHẦN …" heading: named kind first, then number."""
    t = rest.lower()
    if ("đúng" in t and "sai" in t) or "dung sai" in t or "true" in t:
        return "tf"
    if any(k in t for k in ("trả lời ngắn", "tra loi ngan", "điền", "short")):
        return "short"
    if any(k in t for k in ("trắc nghiệm", "trac nghiem", "nhiều lựa chọn", "abcd", "multiple")):
        return "mcq"
    return _PART_QTYPE[re.sub(r"\s|–", lambda m: "-" if m.group() == "–" else "", number.lower())]
TF_STATEMENTS = 4


def _is_red(run) -> bool:
    try:
        rgb = run.font.color.rgb if run.font.color is not None and run.font.color.type is not None else None
    except (AttributeError, ValueError):
        rgb = None
    if rgb is None:
        return False
    r, g, b = rgb[0], rgb[1], rgb[2]
    return r >= 0xB0 and g <= 0x60 and b <= 0x60


# Superscript / subscript text kept as Unicode, one character for one so the
# offsets the format-keeping mixer uses don't move: 130⁰, x², H₂O.
_SUP = str.maketrans("0123456789+-=()n", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁿ")
_SUB = str.maketrans("0123456789+-=()", "₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎")


def _paragraph_chars(paragraph) -> tuple[str, list[bool]]:
    """Paragraph text plus, per character, whether it is red or underlined."""
    text, marked = [], []
    for run in paragraph.runs:
        flag = bool(run.font.underline) or _is_red(run)
        t = run.text
        if run.font.superscript:
            t = t.translate(_SUP)
        elif run.font.subscript:
            t = t.translate(_SUB)
        text.append(t)
        marked.extend([flag] * len(t))
    return "".join(text), marked


# What a lost-content warning says (frontend LOSSY / keep-format filter match it)
LOST = "bị mất"
TOKEN_IDS = re.compile(r"\[\[ct:([0-9a-f]{8,40})\]\]")

_OBJECT_TAGS = ("object", "drawing", "pict")


def _paragraph_objects(paragraph, doc_part) -> list[tuple[int, dict]]:
    """Formulas/pictures of a paragraph as (offset in _paragraph_chars' text,
    asset), in reading order. Offsets count like python-docx's run.text."""
    from app.services.rich_objects import collect

    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
    MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
    out, pos = [], 0
    for child in paragraph._p.iterchildren():
        if child.tag == W + "r":
            run_len = 0
            for c in child.iterchildren():
                if c.tag == W + "t":
                    run_len += len(c.text or "")
                elif c.tag in (W + "tab", W + "br", W + "cr", W + "noBreakHyphen", W + "ptab"):
                    run_len += 1
                el = None
                if c.tag in tuple(W + t for t in _OBJECT_TAGS):
                    el = c
                elif c.tag == MC + "AlternateContent":   # a newer picture, with a VML fallback
                    el = next(iter(c.iterfind(f"{MC}Choice/{W}drawing")), None)
                if el is not None:
                    a = collect(el, doc_part)
                    if a is not None:
                        out.append((pos + run_len, a))
            pos += run_len
        elif child.tag in (M + "oMath", M + "oMathPara"):
            for om in ([child] if child.tag == M + "oMath" else child.iterfind(M + "oMath")):
                a = collect(om, doc_part)
                if a is not None:
                    out.append((pos, a))
    return out


def _rich(text: str, objs: list[tuple[int, dict]], start: int, end: int) -> str:
    """text[start:end] with each formula/picture at its place as a token
    (one sitting right at `end` belongs to this piece)."""
    from app.services.rich_objects import token

    if not any(start <= at <= end for at, _ in objs):
        return text[start:end].strip()          # unchanged when there's nothing to place
    pieces, pos = [], start
    for at, a in objs:
        if start <= at <= end:
            pieces.append(text[pos:at])
            pieces.append(f" {token(a['id'])} " if pieces[-1] and not pieces[-1].endswith((" ", "\t")) else f"{token(a['id'])} ")
            pos = at
    pieces.append(text[pos:end])
    return re.sub(r"[ \t]+", " ", "".join(pieces)).strip()


def _parse_tf_key(raw: str) -> list[bool] | None:
    """"Đ S Đ S", "ĐSĐS", "Đúng, Sai, …", "T F T F" → [True, False, True, False]."""
    s = raw.strip().upper().replace("ĐÚNG", "Đ").replace("SAI", "S").replace("TRUE", "T").replace("FALSE", "F")
    letters = [c for c in s if c in "ĐDSTF"]
    if len(letters) != TF_STATEMENTS:
        return None
    return [c in "ĐDT" for c in letters]


def parse_docx(data: bytes, allow_unanswered: bool = False) -> tuple[list[ParsedQuestion], list[str]]:
    """Questions of a Word đề + warnings. allow_unanswered: keep trắc nghiệm
    questions with no answer marked (answer = -1) instead of skipping them, so
    the teacher can pick the answers on the web."""
    questions, warnings, _ = parse_docx_layout(data, allow_unanswered)
    return questions, warnings


def parse_docx_layout(data: bytes, allow_unanswered: bool = False) -> tuple[list[ParsedQuestion], list[str], dict]:
    """parse_docx plus where things are in the document body (element indexes):
    {"prefix_end": first element that belongs to a question/part/group line,
     "parts": [(index, qtype)], "footer": index of a "HẾT" line or None}.
    Each ParsedQuestion.src holds its own elements and option offsets."""
    from docx import Document   # imported lazily so the rest of the app works without python-docx
    from docx.text.paragraph import Paragraph

    from app.services.rich_objects import token

    try:
        doc = Document(io.BytesIO(data))
    except Exception:
        return [], ["Không đọc được file Word. File phải là .docx (Word 2007 trở lên), không phải .doc"], {}

    questions: list[ParsedQuestion] = []
    warnings:  list[str] = []
    layout: dict = {"prefix_end": None, "parts": [], "footer": None}
    shuffle_options = True
    qtype = "mcq"
    counters = {"mcq": 0, "tf": 0, "short": 0}
    group: dict | None = None     # {"id", "mode", "fixed"} of the current "<gN>" line
    n_groups = 0

    # {"n", "qtype", "content": [..], "options": [..], "marked": set, "key": str|None, "image", "shuffle",
    #  "group", "src": {...}}
    cur: dict | None = None

    def boundary(bi: int) -> None:
        if layout["prefix_end"] is None:
            layout["prefix_end"] = bi

    def add(q: ParsedQuestion, g: dict | None) -> None:
        if g is not None:
            q.group, q.group_mode, q.group_fixed = g["id"], g["mode"], g["fixed"]
        q.src_index = len(questions)
        q.src = cur["src"] if cur else None
        if cur:
            used = set()
            for t in [q.content] + [o.get("text", "") for o in q.options]:
                used.update(TOKEN_IDS.findall(t))
            q.assets = {i: cur["assets"][i] for i in used if i in cur["assets"]}
        questions.append(q)

    def flush() -> None:
        nonlocal cur
        if cur is None:
            return
        kind, content = cur["qtype"], "\n".join(cur["content"]).strip()
        # …and a question with no options/statements but a numeric "Đáp án:"
        # is a trả lời ngắn one (a Bộ GD "PHẦN III" heading read as Đúng/Sai).
        if kind != "short" and not cur["options"] and cur["key"] and normalize_short_answer(cur["key"]):
            kind = "short"
        cur["src"]["qtype"] = kind
        where = f"Phần {PART_OF[kind]} câu {cur['n']}"
        label = f"{where} \"{_short(content)}\""
        if cur["table"]:
            warnings.append(f"{label}: có bảng, bảng {LOST} khi lưu vào ngân hàng hoặc trộn dạng chữ "
                            "(Trộn nhanh từ 1 file Word giữ được bảng)")
        if cur["lost"]:
            warnings.append(f"{label}: có hình ảnh hoặc công thức không đọc được, phần đó {LOST}")
        opts, marked, key = cur["options"], cur["marked"], cur["key"]

        if not content and not cur["image"]:
            warnings.append(f"{where}: thiếu nội dung câu hỏi, bỏ qua")
        elif kind == "mcq":
            if key and not marked and len(key) == 1 and key.upper() in LETTERS[:len(opts)]:
                marked = {LETTERS.index(key.upper())}
            if len(opts) < 2:
                warnings.append(f"{label}: không tìm thấy đáp án A./B./…, bỏ qua")
            elif len(marked) > 1:
                warnings.append(f"{label}: có nhiều hơn 1 đáp án được đánh dấu, bỏ qua")
            elif not marked and not allow_unanswered:
                warnings.append(f"{label}: chưa đánh dấu đáp án đúng (tô đỏ, gạch chân hoặc dòng \"Đáp án: B\"), bỏ qua")
            else:
                add(ParsedQuestion(
                    content=content, options=opts, answer=next(iter(marked)) if marked else -1,
                    shuffle_options=cur["shuffle"] and not _is_true_false(opts), where=where), cur["group"])
        elif kind == "tf":
            if len(opts) != TF_STATEMENTS:
                warnings.append(f"{label}: câu Đúng/Sai phải có đúng 4 ý a) b) c) d) (tìm thấy {len(opts)}), bỏ qua")
            elif key and _parse_tf_key(key) is None:
                warnings.append(f"{label}: dòng \"Đáp án: {key}\" không hợp lệ (cần 4 ký tự Đ/S, vd \"Đ S Đ S\"), bỏ qua")
            else:
                truth = _parse_tf_key(key) if key else [i in marked for i in range(TF_STATEMENTS)]
                if not key and not marked:
                    warnings.append(f"{label}: không có ý nào được đánh dấu Đúng nên hiểu là cả 4 ý Sai, hãy kiểm tra lại")
                add(ParsedQuestion(
                    content=content, qtype="tf", shuffle_options=False, where=where,
                    options=[{**o, "correct": t} for o, t in zip(opts, truth)]), cur["group"])
        else:
            ans = normalize_short_answer(key) if key else None
            if key is None:
                warnings.append(f"{label}: thiếu dòng \"Đáp án: …\", bỏ qua")
            elif ans is None:
                warnings.append(f"{label}: đáp án \"{key}\" không tô được trên phiếu (tối đa 4 ký tự: dấu -, chữ số, dấu phẩy), bỏ qua")
            else:
                add(ParsedQuestion(content=content, qtype="short", answer_text=ans, shuffle_options=False,
                                   where=where), cur["group"])
        cur = None

    W_P, W_TBL = qn_w("p"), qn_w("tbl")
    for bi, el in enumerate(doc.element.body.iterchildren()):
        if el.tag == W_TBL:
            if cur is not None:        # a table inside a question stays with it
                cur["src"]["elems"].append(bi)
                cur["table"] = True
            continue
        if el.tag != W_P:
            continue                   # sectPr, bookmarks…
        p = Paragraph(el, doc)
        text, marked = _paragraph_chars(p)
        # pictures, MathType/Equation objects (w:object) and Word equations (m:oMath) —
        # none of them are text, so they'd silently vanish; flag them instead
        has_image = bool(p._p.xpath(".//w:drawing | .//w:pict | .//w:object | .//m:oMath"))
        # …kept as "[[ct:<id>]]" tokens at their place in the text (rich_objects.py)
        objs = _paragraph_objects(p, doc.part) if has_image else []
        # a floating figure (wp:anchor) has no real place in the line: it
        # illustrates the câu ("Trong hình bên…"), so it goes with its content
        floats = [a for _, a in objs if "<wp:anchor" in a["xml"] or ":anchor " in a["xml"][:400]]
        objs = [(at, a) for at, a in objs if a not in floats]
        float_line = " ".join(token(a["id"]) for a in floats)

        def keep_objects() -> None:     # this line's formulas belong to the current question
            cur["assets"].update({a["id"]: a for _, a in objs})
            cur["assets"].update({a["id"]: a for a in floats})
            if float_line:
                cur["content"].append(float_line)
            if has_image and not objs and not floats:
                cur["lost"] = True

        if not text.strip():
            if cur is not None:
                keep_objects()
                cur["src"]["elems"].append(bi)
                if has_image:
                    cur["image"] = True
                    line = _rich(text, objs, 0, len(text)) if objs else ""
                    if cur["options"]:   # a picture under the options belongs to the last one
                        cur["src"]["extra"].setdefault(len(cur["options"]) - 1, []).append(bi)
                        if line:
                            cur["options"][-1]["text"] = (cur["options"][-1]["text"] + " " + line).strip()
                    elif line:
                        cur["content"].append(line)
            continue

        if _DOCX_END.match(text):
            flush()
            layout["footer"] = bi
            continue

        m_part = _DOCX_PART.match(text)
        if m_part:
            flush()
            boundary(bi)
            layout["footer"] = None
            qtype = _part_qtype(m_part.group(1), m_part.group(2))
            layout["parts"].append((bi, qtype))
            group, shuffle_options = None, True     # groups don't span parts
            continue

        m_group = _DOCX_GROUP.match(text)
        if m_group:
            flush()
            boundary(bi)
            n_groups += 1
            mode = int(m_group.group(2))
            group = {"id": n_groups, "mode": mode, "fixed": bool(m_group.group(1))}
            shuffle_options = mode in (2, 3)
            continue

        m_q = _DOCX_QUESTION.match(text)
        if m_q:
            flush()
            boundary(bi)
            layout["footer"] = None
            counters[qtype] += 1
            first = _rich(text, objs, m_q.end(), len(text))
            cur = {"n": counters[qtype], "qtype": qtype, "content": [f"{first} {float_line}".strip()],
                   "options": [], "marked": set(), "key": None, "image": has_image, "shuffle": shuffle_options,
                   "group": group, "table": False, "lost": has_image and not objs and not floats,
                   "assets": {**{a["id"]: a for _, a in objs}, **{a["id"]: a for a in floats}},
                   "src": {"qline": bi, "number_end": m_q.end(), "number": text[:m_q.end()],
                           "elems": [bi], "slots": [], "extra": {}, "drop": []}}
            continue

        if cur is None:
            continue   # title / instructions / section headings outside questions

        if cur["options"] and _DOCX_SECTION.match(text):
            flush()    # "II/ CĂN THỨC" after the last option: a new section, not option text
            continue

        cur["image"] = cur["image"] or has_image
        keep_objects()
        cur["src"]["elems"].append(bi)

        m_ans = _DOCX_ANSWER.match(text)
        if m_ans:
            cur["key"] = m_ans.group(1).strip()
            cur["src"]["drop"].append(bi)
            continue

        if cur["qtype"] == "short":
            cur["content"].append(_rich(text, objs, 0, len(text)))
            continue

        # Under a bare-number heading the kind can be wrong (a Bộ GD "PHẦN II"
        # is Đúng/Sai, a VJU one trắc nghiệm): statements "a)" where options
        # "A." were expected mean this is a Đúng/Sai question.
        if cur["qtype"] == "mcq" and not cur["options"] and not _DOCX_OPTION.search(text) \
                and (m_st := _DOCX_STATEMENT.match(text)) and m_st.group(2) == "a":
            cur["qtype"] = "tf"

        pattern, letters = (_DOCX_OPTION, LETTERS) if cur["qtype"] == "mcq" else (_DOCX_STATEMENT, LETTERS.lower())
        # Only accept markers continuing the A, B, C… sequence, so text like
        # "vitamin A. Then" inside an option isn't mistaken for a new option.
        markers = []
        for m in pattern.finditer(text):
            if letters.index(m.group(2)) == len(cur["options"]) + len(markers):
                markers.append(m)
        if not markers:
            line = _rich(text, objs, 0, len(text))
            if cur["options"]:
                cur["options"][-1]["text"] = (cur["options"][-1]["text"] + " " + line).strip()
                cur["src"]["extra"].setdefault(len(cur["options"]) - 1, []).append(bi)
            else:
                cur["content"].append(line)
            continue

        # Offsets of this line, for the format-keeping mixer: per option the
        # separator before it, its label ("A. "/"#B. ") and its content.
        slot = {"bi": bi, "prefix_end": markers[0].start(), "opts": []}
        for i, m in enumerate(markers):
            end  = markers[i + 1].start() if i + 1 < len(markers) else len(text)
            body = text[m.end():end].strip()
            # The "A." / "a)" marker itself red/underlined (YoungMix rule), or
            # the whole option text if the teacher coloured that instead.
            marker_marked = any(marked[m.start(2): m.start(2) + 2])
            body_marked   = bool(body) and all(marked[k] for k in range(m.end(), end) if not text[k].isspace())
            if marker_marked or body_marked:
                cur["marked"].add(len(cur["options"]))
            if body_marked:     # the option text itself is coloured: the mixer uncolours it
                cur["src"].setdefault("body_marked", []).append(len(cur["options"]))
            slot["opts"].append({"index": len(cur["options"]), "sep": m.start(), "label": m.start(1),
                                 "content": m.end(), "end": end})
            cur["options"].append({"text": _rich(text, objs, m.end(), end), "fixed": bool(m.group(1))})
        cur["src"]["slots"].append(slot)

    flush()
    return questions, warnings, layout


def qn_w(tag: str) -> str:
    return "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}" + tag


def export_docx(questions: list[ParsedQuestion], title: str = "", assets: dict | None = None) -> bytes:
    """Re-importing the exported file gives the same questions. Grouped into
    PHẦN I-II / III / IV (only the parts that have questions); the correct
    answer is marked the way the importer reads it: red + underlined letter
    for trắc nghiệm and the true statements of Đúng/Sai, "Đáp án: …" for
    trả lời ngắn. Formulas/pictures ("[[ct:<id>]]") go back in as the real
    objects from `assets`."""
    from docx import Document
    from docx.shared import Pt, RGBColor

    from app.services.rich_objects import add_rich

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn_w("eastAsia"), "Times New Roman")

    if title:   # same font as the rest, not Word's blue Calibri "Heading 1"
        t = doc.add_paragraph().add_run(title)
        t.bold, t.font.size = True, Pt(14)

    def marker(par, label: str, correct: bool) -> None:
        run = par.add_run(label)
        if correct:
            run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
            run.underline = True

    for kind in QTYPES:
        group = [q for q in questions if q.qtype == kind]
        if not group:
            continue
        doc.add_paragraph().add_run(f"PHẦN {PART_OF[kind]}. {PART_NAME[kind]}").bold = True
        for n, q in enumerate(group, start=1):
            p = doc.add_paragraph()
            p.add_run(f"Câu {n}. ").bold = True
            add_rich(doc, p, q.content, assets)
            if kind == "short":
                doc.add_paragraph(f"Đáp án: {q.answer_text}")
                continue
            for i, o in enumerate(q.options):
                op = doc.add_paragraph()
                if kind == "mcq":
                    marker(op, f"{'#' if o.get('fixed') else ''}{LETTERS[i]}.", i == q.answer)
                else:
                    marker(op, f"{LETTERS[i].lower()})", bool(o.get("correct")))
                op.add_run(" ")
                add_rich(doc, op, o["text"], assets)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
