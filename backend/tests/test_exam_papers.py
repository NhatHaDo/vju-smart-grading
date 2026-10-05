"""
test_exam_papers.py
===================
Trộn đề: mixing rules, the /exam-papers API, Word export, attaching to a
kỳ thi, and the per-mã-đề answer key used for grading.
"""
import io
import zipfile

import pytest

from app.services.exam_mixer import (
    MixError,
    answer_key,
    build_versions,
    check_counts,
    grading_key_mau40,
    usable,
    version_codes,
)
from app.services.question_io import ParsedQuestion, parse_docx


def _mcq(i: int, n_opts: int = 4, fixed_last: bool = False, shuffle: bool = True) -> ParsedQuestion:
    opts = [{"text": f"q{i}-opt{k}", "fixed": fixed_last and k == n_opts - 1} for k in range(n_opts)]
    return ParsedQuestion(content=f"Câu {i}", options=opts, answer=i % n_opts, shuffle_options=shuffle)


def _tf(i: int) -> ParsedQuestion:
    return ParsedQuestion(content=f"TF {i}", qtype="tf", shuffle_options=False,
                          options=[{"text": f"s{k}", "fixed": False, "correct": k % 2 == 0} for k in range(4)])


def _short(i: int, ans: str) -> ParsedQuestion:
    return ParsedQuestion(content=f"Số {i}", qtype="short", answer_text=ans, shuffle_options=False)


# ── Mixing rules ─────────────────────────────────────────────────────────────

def test_versions_same_questions_answers_follow_the_shuffle():
    qs = [_mcq(i) for i in range(10)] + [_tf(0), _tf(1)] + [_short(0, "-1,5"), _short(1, "75")]
    versions = build_versions(qs, ["101", "102", "103"], seed=7)
    for code, snaps in versions.items():
        assert [s.qtype for s in snaps] == ["mcq"] * 10 + ["tf"] * 2 + ["short"] * 2
        assert sorted(s.content for s in snaps) == sorted(q.content for q in qs)   # same questions
        for s in snaps:
            if s.qtype == "mcq":
                src = next(q for q in qs if q.content == s.content)
                # the option marked correct is still the one that was correct
                assert s.options[s.answer]["text"] == src.options[src.answer]["text"]
    orders = {code: [s.content for s in snaps] for code, snaps in versions.items()}
    assert len({tuple(o) for o in orders.values()}) > 1   # orders actually differ


def test_fixed_and_unshuffled_options_stay_put():
    qs = [_mcq(i, fixed_last=True) for i in range(8)] + [_mcq(99, shuffle=False)]
    for snaps in build_versions(qs, [str(c) for c in range(101, 111)], seed=3).values():
        for s in snaps:
            if s.content == "Câu 99":
                assert [o["text"] for o in s.options] == [f"q99-opt{k}" for k in range(4)]
            else:
                assert s.options[3]["text"].endswith("opt3")   # "#D." pinned


def test_no_shuffle_keeps_everything_in_order():
    qs = [_mcq(i) for i in range(5)]
    snaps = build_versions(qs, ["101"], shuffle_questions=False, shuffle_options=False)["101"]
    assert [s.content for s in snaps] == [q.content for q in qs]
    assert [s.answer for s in snaps] == [q.answer for q in qs]


def test_counts_limits_and_codes():
    check_counts({"mcq": 50, "tf": 8, "short": 6}, {"mcq": 40, "tf": 8, "short": 6})
    with pytest.raises(MixError, match="chỉ có 40"):
        check_counts({"mcq": 50}, {"mcq": 41})
    with pytest.raises(MixError, match="chỉ có 3 câu dùng được"):
        check_counts({"tf": 3}, {"tf": 4})
    with pytest.raises(MixError):
        check_counts({"mcq": 5}, {"mcq": 0})
    assert version_codes("101", 4) == ["101", "102", "103", "104"]
    assert version_codes("001", 2) == ["001", "002"]
    with pytest.raises(MixError):
        version_codes("abc", 2)          # the sheet only takes digits
    with pytest.raises(MixError):
        version_codes("101", 0)
    with pytest.raises(MixError):
        version_codes("998", 3)          # 1000 doesn't fit 3 bubbles
    # Đề chỉ để in: any code
    assert version_codes("A", 4, for_sheet=False) == ["A", "B", "C", "D"]
    assert version_codes("Đề 1", 2, for_sheet=False) == ["Đề 1", "Đề 2"]
    assert version_codes("MD09", 2, for_sheet=False) == ["MD09", "MD10"]
    assert version_codes("A1, B2 ;C3", 99, for_sheet=False) == ["A1", "B2", "C3"]
    assert version_codes("201, 305", 1) == ["201", "305"]
    for bad in ("Y", "A, a", "x/y", "Đề", "ABCDEFGHIJK1"):
        with pytest.raises(MixError):
            version_codes(bad, 3, for_sheet=False)


def test_grading_key_uses_sheet_labels():
    qs = [_mcq(0), _mcq(1), _tf(0), _short(0, "-1,5")]
    snaps = build_versions(qs, ["101"], shuffle_questions=False, shuffle_options=False)["101"]
    key = grading_key_mau40(snaps, ["p3a", "p3b"])
    assert key == {
        "trc_nghim_abcd1": "A", "trc_nghim_abcd2": "B",
        "ng_sai_cu1": "Đ", "ng_sai_cu2": "S", "ng_sai_cu3": "Đ", "ng_sai_cu4": "S",
        "p3a": "-1.5",   # the OMR reads Phần III with a "." decimal point
    }
    assert answer_key(snaps) == {"mcq": ["A", "B"], "tf": [["Đ", "S", "Đ", "S"]], "short": ["-1,5"]}


# ── API ──────────────────────────────────────────────────────────────────────

def _bank(client, h, n_mcq=12, n_tf=3, n_short=2, n_mcq5=1) -> int:
    base = "/api/v1/question-bank"
    cid = client.post(f"{base}/categories", json={"name": "Ngân hàng"}, headers=h).json()["id"]
    post = lambda body: client.post(f"{base}/categories/{cid}/questions", json=body, headers=h)  # noqa: E731
    for i in range(n_mcq):
        post({"qtype": "mcq", "content": f"MCQ {i}", "options": [{"text": f"{i}{c}"} for c in "abcd"], "answer": i % 4})
    for i in range(n_mcq5):   # 5 options → can't go on the sheet
        post({"qtype": "mcq", "content": f"MCQ5 {i}", "options": [{"text": c} for c in "abcde"], "answer": 4})
    for i in range(n_tf):
        post({"qtype": "tf", "content": f"TF {i}", "options": [{"text": f"s{k}", "correct": k == i % 4} for k in range(4)]})
    for i in range(n_short):
        post({"qtype": "short", "content": f"Số {i}", "answer_text": str(10 + i)})
    return cid


def _exam(client, h, name="Giữa kì") -> int:
    r = client.post("/api/v1/exams", json={"name": name}, headers=h)
    assert r.status_code == 201, r.text
    return r.json()["id"]


def test_api_from_bank_export_attach_and_grading_key(client):
    h1, h2 = client.headers_for(1), client.headers_for(2)
    cid = _bank(client, h1)
    exam_id = _exam(client, h1)
    body = {"name": "Đề giữa kì", "category_ids": [cid], "counts": {"mcq": 10, "tf": 2, "short": 2},
            "num_versions": 3, "start_code": "101"}

    # Validation: more than available / more than the sheet allows
    assert client.post("/api/v1/exam-papers/from-bank", json={**body, "counts": {"mcq": 13}}, headers=h1).status_code == 422
    assert client.post("/api/v1/exam-papers/from-bank", json={**body, "counts": {"tf": 9}}, headers=h1).status_code == 422
    # Another teacher can't use a category that isn't shared with them
    assert client.post("/api/v1/exam-papers/from-bank", json=body, headers=h2).status_code == 404

    r = client.post("/api/v1/exam-papers/from-bank", json=body, headers=h1)
    assert r.status_code == 201, r.text
    paper = r.json()
    assert [v["code"] for v in paper["versions"]] == ["101", "102", "103"]
    assert paper["counts"] == {"mcq": 10, "tf": 2, "short": 2}
    assert paper["notes"] == ["Đã bỏ 1 câu trắc nghiệm có hơn 4 đáp án (phiếu chỉ có A–D)"]
    assert all(len(v["answer_key"]["mcq"]) == 10 for v in paper["versions"])
    pid = paper["id"]

    # Owner-only
    assert client.get(f"/api/v1/exam-papers/{pid}", headers=h2).status_code == 404
    assert client.get("/api/v1/exam-papers", headers=h2).json() == []

    # Word export: one mã đề, the đáp án, and a zip with everything
    r = client.get(f"/api/v1/exam-papers/{pid}/versions/102/docx", headers=h1)
    assert r.status_code == 200
    from docx import Document
    text = "\n".join(p.text for p in Document(io.BytesIO(r.content)).paragraphs)
    assert "Mã đề: 102" in text and "PHẦN I" in text and "PHẦN III" in text
    assert client.get(f"/api/v1/exam-papers/{pid}/answer-key/docx", headers=h1).status_code == 200
    z = zipfile.ZipFile(io.BytesIO(client.get(f"/api/v1/exam-papers/{pid}/zip", headers=h1).content))
    assert sorted(z.namelist()) == ["Dap an.docx", "Ma de 101.docx", "Ma de 102.docx", "Ma de 103.docx"]

    # Not attached yet → no answer key for the kỳ thi
    assert client.get(f"/api/v1/exam-papers/exam-answer-key/{exam_id}", headers=h1).json()["byMaDe"] == {}

    # Attach, then drop mã đề 103 from the kỳ thi
    r = client.put(f"/api/v1/exam-papers/{pid}", json={"exam_id": exam_id}, headers=h1)
    assert r.json()["exam_name"] == "Giữa kì"
    v103 = next(v for v in paper["versions"] if v["code"] == "103")
    client.put(f"/api/v1/exam-papers/{pid}/versions/{v103['id']}", json={"in_exam": False}, headers=h1)

    key = client.get(f"/api/v1/exam-papers/exam-answer-key/{exam_id}", headers=h1).json()
    assert key["versions"] == ["101", "102"] and key["papers"] == ["Đề giữa kì"]
    k101 = key["byMaDe"]["101"]
    assert sum(1 for k in k101 if k.startswith("trc_nghim_abcd")) == 10
    assert sum(1 for k in k101 if k.startswith("ng_sai_cu")) == 8
    shorts = [k for k in k101 if k.startswith("custom_")]
    assert len(shorts) == 2 and all(k101[k] in ("10", "11") for k in shorts)
    # mã đề 101's key matches its own printed đáp án
    v101 = next(v for v in paper["versions"] if v["code"] == "101")
    assert [k101[f"trc_nghim_abcd{i}"] for i in range(1, 11)] == v101["answer_key"]["mcq"]

    # A second bộ đề in the same kỳ thi can't reuse mã đề 101/102
    r = client.post("/api/v1/exam-papers/from-bank", json={**body, "name": "Đề 2", "exam_id": exam_id}, headers=h1)
    pid2 = r.json()["id"]
    # (created with exam_id: codes clash is checked when attaching via PUT)
    client.put(f"/api/v1/exam-papers/{pid2}", json={"exam_id": None}, headers=h1)
    r = client.put(f"/api/v1/exam-papers/{pid2}", json={"exam_id": exam_id}, headers=h1)
    assert r.status_code == 409 and "101" in r.json()["detail"]

    # Detach / delete
    assert client.put(f"/api/v1/exam-papers/{pid}", json={"exam_id": None}, headers=h1).json()["exam_id"] is None
    assert client.delete(f"/api/v1/exam-papers/{pid}", headers=h1).status_code == 204
    assert client.get(f"/api/v1/exam-papers/{pid}", headers=h1).status_code == 404


def test_api_from_file(client, tmp_path):
    from docx import Document
    doc = Document()
    doc.add_paragraph("PHẦN I")
    for i in range(3):
        doc.add_paragraph(f"Câu {i + 1}. Hỏi {i}")
        doc.add_paragraph("A. a\tB. b\tC. c\tD. d")
        doc.add_paragraph(f"Đáp án: {'ABC'[i]}")
    doc.add_paragraph("PHẦN III")
    doc.add_paragraph("Câu 1. Bao nhiêu?")
    doc.add_paragraph("Đáp án: 0,5")
    buf = io.BytesIO()
    doc.save(buf)

    h = client.headers_for(1)
    r = client.post("/api/v1/exam-papers/from-file", headers=h,
                    data={"name": "Đề từ file", "num_versions": "2", "start_code": "201"},
                    files={"file": ("de.docx", buf.getvalue(), "application/octet-stream")})
    assert r.status_code == 201, r.text
    p = r.json()
    assert p["source"] == "file" and p["counts"] == {"mcq": 3, "tf": 0, "short": 1}
    assert [v["code"] for v in p["versions"]] == ["201", "202"]
    assert all(len(v["answer_key"]["mcq"]) == 3 and v["answer_key"]["short"] == ["0,5"] for v in p["versions"])

    r = client.post("/api/v1/exam-papers/from-file", headers=h, data={"name": "x"},
                    files={"file": ("de.pdf", b"%PDF", "application/pdf")})
    assert r.status_code == 400


def test_api_from_big_moodle_file_picks_random_subset(client):
    """A 131-câu Moodle export: too many for the sheet as-is, so the teacher
    reads the file first and picks how many to take."""
    txt = "\n\n".join(f"Câu hỏi số {i}\nA. a{i}\nB. b{i}\nC. c{i}\nD. d{i}\nANSWER: B" for i in range(131))
    h = client.headers_for(1)
    f = lambda: {"file": ("big.txt", txt.encode(), "text/plain")}

    r = client.post("/api/v1/exam-papers/parse-file", headers=h, files=f())
    assert r.status_code == 200, r.text
    assert r.json()["available"] == {"mcq": 131, "tf": 0, "short": 0}
    assert r.json()["limits"]["mcq"] == 40

    r = client.post("/api/v1/exam-papers/from-file", headers=h, data={"name": "Tất cả"}, files=f())
    assert r.status_code == 422 and "40" in r.json()["detail"]

    r = client.post("/api/v1/exam-papers/from-file", headers=h, files=f(),
                    data={"name": "40 câu", "num_versions": "3", "count_mcq": "40"})
    assert r.status_code == 201, r.text
    p = r.json()
    assert p["counts"] == {"mcq": 40, "tf": 0, "short": 0}
    assert any("40/131" in n for n in p["notes"])
    full = client.get(f"/api/v1/exam-papers/{p['id']}", headers=h).json()
    stems = [sorted(q["content"] for q in v["questions"]) for v in full["versions"]]
    assert len(stems[0]) == 40 and all(s == stems[0] for s in stems)   # same 40 in every mã đề

    r = client.post("/api/v1/exam-papers/from-file", headers=h, files=f(),
                    data={"name": "x", "count_mcq": "10", "count_tf": "1"})
    assert r.status_code == 422   # no Đúng/Sai in the file


def test_api_print_only_mix_no_sheet_limits(client):
    """Đề chỉ để in (không chấm bằng phiếu Mẫu 40): every question, 5-option
    questions kept — but it can't be attached to a kỳ thi."""
    qs = [f"Câu hỏi số {i}\nA. a\nB. b\nC. c\nD. d\nANSWER: A" for i in range(130)]
    qs.append("Câu năm đáp án\nA. a\nB. b\nC. c\nD. d\nE. e\nANSWER: E")
    txt = "\n\n".join(qs).encode()
    h = client.headers_for(1)
    f = lambda: {"file": ("in.txt", txt, "text/plain")}

    r = client.post("/api/v1/exam-papers/parse-file", headers=h, files=f())
    assert r.json()["available"]["mcq"] == 130 and r.json()["available_all"]["mcq"] == 131

    r = client.post("/api/v1/exam-papers/from-file", headers=h, files=f(),
                    data={"name": "Đề in", "num_versions": "2", "start_code": "A", "for_sheet": "false"})
    assert r.status_code == 201, r.text
    p = r.json()
    assert p["counts"]["mcq"] == 131 and p["gradable"] is False and "131" in p["sheet_problem"]
    assert all(len(v["answer_key"]["mcq"]) == 131 for v in p["versions"])
    assert [v["code"] for v in p["versions"]] == ["A", "B"]
    assert client.get(f"/api/v1/exam-papers/{p['id']}/versions/B/docx", headers=h).status_code == 200
    assert client.get(f"/api/v1/exam-papers/{p['id']}/zip", headers=h).status_code == 200

    exam_id = _exam(client, h, "Cuối kì")
    r = client.put(f"/api/v1/exam-papers/{p['id']}", headers=h, json={"exam_id": exam_id})
    assert r.status_code == 422 and "Mẫu 40" in r.json()["detail"]
    r = client.post("/api/v1/exam-papers/from-file", headers=h, files=f(),
                    data={"name": "x", "for_sheet": "false", "exam_id": str(exam_id)})
    assert r.status_code == 422

    # ≤ 40 câu but mã đề "A"/"B": still can't be bubbled on the sheet
    four = "\n\n".join(qs[:10]).encode()   # no 5-option question
    r = client.post("/api/v1/exam-papers/from-file", headers=h, files={"file": ("in.txt", four, "text/plain")},
                    data={"name": "chữ", "for_sheet": "false", "start_code": "A"})
    assert r.status_code == 201 and r.json()["gradable"] is False and "mã đề" in r.json()["sheet_problem"]

    # Chọn 40 câu, vẫn để in: fits the sheet only if no 5-option question got picked
    r = client.post("/api/v1/exam-papers/from-file", headers=h, files=f(),
                    data={"name": "40 câu", "for_sheet": "true", "count_mcq": "40"})
    assert r.status_code == 201 and r.json()["gradable"] is True


def test_api_from_several_files_pooled(client):
    """Nhiều file (vd mỗi chương 1 file) → one pool of questions."""
    from pathlib import Path
    samples = Path(__file__).parent.parent.parent / "frontend" / "public" / "samples"
    moodle = (samples / "de-mau-moodle.txt").read_bytes()
    simple = (samples / "de-trac-nghiem-don-gian.docx").read_bytes()
    grouped = (samples / "de-mau-trac-nghiem-nhom.docx").read_bytes()
    bad = b"Screwdriver\nANSWER: \n"
    h = client.headers_for(1)
    files = [("file", ("chuong1.txt", moodle, "text/plain")),
             ("file", ("chuong2.docx", simple, "application/octet-stream")),
             ("file", ("nhom.docx", grouped, "application/octet-stream")),
             ("file", ("loi.txt", bad, "text/plain"))]

    r = client.post("/api/v1/exam-papers/parse-file", headers=h, files=files)
    assert r.status_code == 200, r.text
    info = r.json()
    n_grouped = len(parse_docx(grouped)[0])
    # the 5 câu of the Moodle sample are also in the simple Word sample: pooled once
    assert info["available"]["mcq"] == 5 + 10 + n_grouped - 5
    assert len(info["duplicate_list"]) == 5 and all(d.startswith("chuong2.docx: ") for d in info["duplicate_list"])
    assert info["warnings"] and info["warnings"][0].startswith("loi.txt: ")

    r = client.post("/api/v1/exam-papers/from-file", headers=h, files=files,
                    data={"name": "Gộp", "for_sheet": "false", "num_versions": "2"})
    assert r.status_code == 201, r.text
    p = r.json()
    assert p["counts"]["mcq"] == 5 + 10 + n_grouped - 5
    assert p["settings"]["file_name"] == "chuong1.txt, chuong2.docx, nhom.docx, loi.txt"


def test_deleting_exam_detaches_its_papers(client):
    h = client.headers_for(1)
    _bank(client, h)
    exam_id = _exam(client, h, "Sẽ xóa")
    cats = client.get("/api/v1/question-bank/categories", headers=h).json()
    r = client.post("/api/v1/exam-papers/from-bank", headers=h, json={
        "name": "Gắn rồi xóa", "category_ids": [cats[0]["id"]], "counts": {"mcq": 5},
        "num_versions": 2, "exam_id": exam_id})
    assert r.status_code == 201, r.text
    paper_id = r.json()["id"]
    assert client.delete(f"/api/v1/exams/{exam_id}", headers=h).status_code in (200, 204)
    p = client.get(f"/api/v1/exam-papers/{paper_id}", headers=h).json()
    assert p["exam_id"] is None and p["exam_name"] is None


def test_parse_then_mix_real_sample_file():
    """The downloadable sample đề (3 phần) imports cleanly and mixes."""
    from pathlib import Path
    # The same file teachers download from the Import dialog ("Tải file mẫu")
    sample = Path(__file__).parent.parent.parent / "frontend" / "public" / "samples" / "de-mau-3-phan.docx"
    qs, warnings = parse_docx(sample.read_bytes())
    assert warnings == []
    versions = build_versions(qs, ["101", "102"], seed=1)
    assert all(len(v) == len(qs) for v in versions.values())


# ── YoungMix groups <g0>–<g3>, <#gN> ─────────────────────────────────────────

@pytest.mark.parametrize("name,expect", [
    ("de-trac-nghiem-don-gian.docx", {"mcq": 10, "tf": 0, "short": 0}),
    ("de-chuan-mau-40.docx",         {"mcq": 40, "tf": 8, "short": 6}),   # exactly fills the sheet
    ("de-mau-trac-nghiem-nhom.docx", None),
    ("de-mau-moodle.txt",            {"mcq": 5, "tf": 0, "short": 0}),
])
def test_downloadable_samples_parse_cleanly(name, expect):
    """Every file mẫu on the Trộn đề page reads with no warning and fits the sheet."""
    from pathlib import Path
    from app.services.exam_mixer import sheet_problem
    from app.services.question_io import QTYPES, parse_aiken
    data = (Path(__file__).parent.parent.parent / "frontend" / "public" / "samples" / name).read_bytes()
    qs, warnings = parse_aiken(data.decode("utf-8")) if name.endswith(".txt") else parse_docx(data)
    assert warnings == []
    if expect:
        assert {qt: sum(q.qtype == qt for q in qs) for qt in QTYPES} == expect
    assert all(usable(q) for q in qs)
    v = build_versions(qs, ["101"], seed=3)["101"]
    assert sheet_problem(v, "101") is None


def _grouped_docx() -> bytes:
    from docx import Document
    doc = Document()
    n = 0
    def q(tag: str):
        nonlocal n
        n += 1
        doc.add_paragraph(f"Câu {n}. {tag}-{n}")
        doc.add_paragraph(f"A. {n}a\tB. {n}b\tC. {n}c\tD. {n}d")
        doc.add_paragraph("Đáp án: A")
    doc.add_paragraph("<#g0>")          # e.g. a listening section: fixed, untouched
    for _ in range(3): q("g0")
    doc.add_paragraph("<g1>")
    for _ in range(4): q("g1")
    doc.add_paragraph("<g2>")
    for _ in range(3): q("g2")
    doc.add_paragraph("<g3>")
    for _ in range(4): q("g3")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_parse_groups():
    qs, warnings = parse_docx(_grouped_docx())
    assert warnings == [] and len(qs) == 14
    assert [(q.group_mode, q.group_fixed) for q in qs[:4]] == [(0, True)] * 3 + [(1, False)]
    assert [q.shuffle_options for q in qs] == [False] * 7 + [True] * 7
    assert len({q.group for q in qs}) == 4


def test_groups_rules_in_every_version():
    qs, _ = parse_docx(_grouped_docx())
    versions = build_versions(qs, [str(c) for c in range(101, 121)], seed=11)
    orders_g1, orders_g3, block_orders = set(), set(), set()
    for snaps in versions.values():
        tags = [s.content.split("-")[0] for s in snaps]
        # <#g0> pinned first, questions and options untouched
        assert [s.content for s in snaps[:3]] == ["g0-1", "g0-2", "g0-3"]
        assert all([o["text"] for o in s.options] == [f"{s.content[3:]}{c}" for c in "abcd"] for s in snaps[:3])
        # every group stays contiguous
        for g in ("g1", "g2", "g3"):
            idx = [i for i, t in enumerate(tags) if t == g]
            assert idx == list(range(idx[0], idx[0] + len(idx)))
        block_orders.add(tuple(dict.fromkeys(tags[3:])))
        g1 = [s for s in snaps if s.content.startswith("g1")]
        g2 = [s for s in snaps if s.content.startswith("g2")]
        g3 = [s for s in snaps if s.content.startswith("g3")]
        orders_g1.add(tuple(s.content for s in g1))
        orders_g3.add(tuple(s.content for s in g3))
        # <g1>: options untouched · <g2>: question order untouched
        assert all(s.answer == 0 for s in g1)
        assert [s.content for s in g2] == ["g2-8", "g2-9", "g2-10"]
    assert len(orders_g1) > 1 and len(orders_g3) > 1   # question order shuffled inside g1 / g3
    assert len(block_orders) > 1                        # unpinned groups swap places
    # <g2>/<g3> options do get shuffled somewhere
    assert any(s.answer != 0 for snaps in versions.values() for s in snaps if s.content[:2] in ("g2", "g3"))


def test_nothing_read_message_names_the_real_reason():
    """A practice file with no marked answers and MathType options: the error
    says "no answer marked" first and counts questions per reason, instead of
    quoting the first warning (a formula one) with a warning total."""
    from app.api.v1.routes.exam_papers import _why_nothing_read
    warnings = [
        'Phần I-II câu 1 "Căn bậc hai": có bảng, bảng bị mất khi lưu vào ngân hàng hoặc trộn dạng chữ',
        'Phần I-II câu 1 "Căn bậc hai": chưa đánh dấu đáp án đúng (tô đỏ, gạch chân hoặc dòng "Đáp án: B"), bỏ qua',
        'Phần I-II câu 2 "Tính": chưa đánh dấu đáp án đúng (tô đỏ, gạch chân hoặc dòng "Đáp án: B"), bỏ qua',
    ]
    msg = _why_nothing_read(warnings)
    assert msg.index("2 câu chưa đánh dấu đáp án đúng") < msg.index("1 câu có bảng")
    assert "3 câu" not in msg


def test_mix_from_file_drops_exact_repeats(client):
    """The same câu twice in a file: listed, and left out so a mã đề never has it twice."""
    h = client.headers_for(1)
    q = "Which is the fastest cache?\nA) L1\nB) L2\nC) L3\nD) L4\nANSWER: A\n"
    other = "Which unit is a byte?\nA) 8 bits\nB) 4 bits\nC) 2 bits\nD) 1 bit\nANSWER: A\n"
    data = (q + other + q).encode()
    f = lambda: {"file": ("de.txt", data, "text/plain")}
    info = client.post("/api/v1/exam-papers/parse-file", headers=h, files=f()).json()
    assert info["available_all"]["mcq"] == 2
    assert len(info["duplicate_list"]) == 1 and "trùng Câu 1" in info["duplicate_list"][0]
    r = client.post("/api/v1/exam-papers/from-file", headers=h, files=f(),
                    data={"name": "Trùng", "num_versions": "2", "for_sheet": "false"})
    assert r.status_code == 201, r.text
    p = r.json()
    assert p["counts"]["mcq"] == 2
    assert any("bỏ 1 câu trùng" in n for n in p["notes"])
    for v in p["versions"]:
        assert len(v["answer_key"]["mcq"]) == 2


def _install_sheets(client, tmp_path, monkeypatch):
    """The shared Mẫu 40 + Bộ GD sheets in the test DB; returns (db, mau40 id, bgd id)."""
    from app.database import get_db
    from app.services import shared_templates as st
    for attr, fname in (("MAU40_DEST_TPL", "shared_40tn_dungsai.template.json"),
                        ("MAU40_DEST_AREAS", "shared_40tn_dungsai.areas.json"),
                        ("BGD_DEST_TPL", "shared_bgd_40tn.template.json"),
                        ("BGD_DEST_AREAS", "shared_bgd_40tn.areas.json")):
        monkeypatch.setattr(st, attr, tmp_path / fname)
    monkeypatch.setattr(st, "DEST_DIR", tmp_path)
    db = next(client.app.dependency_overrides[get_db]())
    st.ensure_shared_templates(db)
    return db, st.find_mau40(db).id, st.find_bgd(db).id


def _bank_with_mcq(client, h, n=3):
    cat = client.post("/api/v1/question-bank/categories", json={"name": "Phiếu"}, headers=h).json()
    for i in range(n):
        client.post(f"/api/v1/question-bank/categories/{cat['id']}/questions",
                    json={"content": f"Câu {i}?", "options": [{"text": "a"}, {"text": "b"}, {"text": "c"}, {"text": "d"}], "answer": 1},
                    headers=h)
    return cat


def test_bo_de_remembers_its_answer_sheet(client, tmp_path, monkeypatch):
    """anh Tú: "cái này a chỉ chọn đc mỗi mẫu phiếu 40 câu thôi à" — a bộ đề
    can be mixed for the Bộ GD sheet too; the choice comes back on the bộ đề
    and on the kỳ thi's answer key."""
    _, mau40, bgd = _install_sheets(client, tmp_path, monkeypatch)
    h = client.headers_for(1)
    cat = _bank_with_mcq(client, h)
    exam = client.post("/api/v1/exams", json={"name": "Kỳ thi phiếu Bộ GD", "subject": "Tin"}, headers=h).json()
    body = {"name": "Đề BGD", "category_ids": [cat["id"]], "counts": {"mcq": 3, "tf": 0, "short": 0},
            "num_versions": 2, "start_code": "101", "exam_id": exam["id"], "sheet": bgd}
    paper = client.post("/api/v1/exam-papers/from-bank", json=body, headers=h).json()
    assert paper["sheet"] == bgd and paper["sheet_name"] == "Phiếu Bộ GD"
    key = client.get(f"/api/v1/exam-papers/exam-answer-key/{exam['id']}", headers=h).json()
    assert key["sheets"] == [bgd] and key["versions"] == ["101", "102"]
    assert set(key["byMaDe"]["101"]) == {"trc_nghim_abcd1", "trc_nghim_abcd2", "trc_nghim_abcd3"}
    # grading with the other sheet gets nothing from this bộ đề
    assert client.get(f"/api/v1/exam-papers/exam-answer-key/{exam['id']}?sheet={mau40}", headers=h).json()["versions"] == []
    assert client.get(f"/api/v1/exam-papers/exam-answer-key/{exam['id']}?sheet={bgd}", headers=h).json()["versions"] == ["101", "102"]
    # the names sent before 2026-10-05 still work; the Bộ GD sheet has 3 mã đề columns
    assert client.post("/api/v1/exam-papers/from-bank", json={**body, "name": "Đề VJU", "exam_id": None, "sheet": "mau40"},
                       headers=h).json()["sheet"] == mau40
    r = client.post("/api/v1/exam-papers/from-bank", json={**body, "exam_id": None, "start_code": "1001"}, headers=h)
    assert r.status_code == 422 and "tối đa 3 chữ số" in r.json()["detail"]
    # default stays Mẫu 40; chỉ in đề has no sheet
    del body["sheet"]
    assert client.post("/api/v1/exam-papers/from-bank", json={**body, "name": "Mặc định", "exam_id": None},
                       headers=h).json()["sheet"] == mau40
    assert client.post("/api/v1/exam-papers/from-bank", json={**body, "name": "In", "exam_id": None, "for_sheet": False},
                       headers=h).json()["sheet"] is None


def test_any_answer_sheet_can_be_picked(client, tmp_path, monkeypatch):
    """"cái này có nhiều mẫu phiếu lắm mà": the teacher's own custom template
    is offered too, and what a bộ đề may hold comes from the sheet itself —
    here a copy of the Bộ GD sheet without its Mã đề box (only 1 mã đề)."""
    import json as _json
    from app.models.template import Template
    from app.services import shared_templates as st
    db, mau40, bgd = _install_sheets(client, tmp_path, monkeypatch)
    areas = [a for a in _json.loads(st.BGD_SRC_AREAS.read_text(encoding="utf-8")) if a.get("blockName") != "made"]
    (tmp_path / "own.areas.json").write_text(_json.dumps(areas), encoding="utf-8")
    (tmp_path / "own.template.json").write_text(st.BGD_SRC_TPL.read_text(encoding="utf-8"), encoding="utf-8")
    own = Template(name="Phiếu của tôi", type="custom", version="1.0", owner_user_id=1, is_default=False,
                   file_path=str(tmp_path / "own.template.json"), areas_path=str(tmp_path / "own.areas.json"))
    db.add(own)
    db.commit()

    h = client.headers_for(1)
    sheets = client.get("/api/v1/exam-papers/sheets", headers=h).json()
    assert [s["id"] for s in sheets] == [mau40, bgd, own.id]
    assert sheets[0]["limits"] == {"mcq": 40, "tf": 8, "short": 6} and sheets[0]["code_digits"] == 4
    assert sheets[1]["code_digits"] == 3 and sheets[2]["code_digits"] is None and sheets[2]["mcq_options"] == 4
    # another teacher doesn't see it, and can't mix for it
    assert [s["id"] for s in client.get("/api/v1/exam-papers/sheets", headers=client.headers_for(2)).json()] == [mau40, bgd]

    cat = _bank_with_mcq(client, h)
    body = {"name": "Phiếu riêng", "category_ids": [cat["id"]], "counts": {"mcq": 3, "tf": 0, "short": 0},
            "num_versions": 2, "start_code": "101", "sheet": own.id}
    r = client.post("/api/v1/exam-papers/from-bank", json=body, headers=h)
    assert r.status_code == 422 and "chỉ trộn được 1 mã đề" in r.json()["detail"]
    paper = client.post("/api/v1/exam-papers/from-bank", json={**body, "num_versions": 1}, headers=h).json()
    assert paper["sheet"] == own.id and paper["sheet_name"] == "Phiếu của tôi" and paper["gradable"]
    cat2 = _bank_with_mcq(client, client.headers_for(2))
    r = client.post("/api/v1/exam-papers/from-bank", json={**body, "category_ids": [cat2["id"]], "num_versions": 1},
                    headers=client.headers_for(2))
    assert r.status_code == 422


def test_grading_key_follows_the_sheet():
    """A sheet with A–E trắc nghiệm takes 5-option câu; the đáp án goes into
    that sheet's own field keys."""
    from app.services.exam_mixer import Snapshot, grading_key, sheet_problem
    from app.services.sheet_layouts import SheetLayout
    lay = SheetLayout(id=9, name="A–E", mcq=["q1", "q2"], mcq_options=5, tf=["t1", "t2", "t3", "t4"], short=["s1"],
                      code_digits=2)
    snaps = [Snapshot("mcq", "x", [{"text": str(i)} for i in range(5)], answer=4),
             Snapshot("tf", "y", [{"text": "a", "correct": True}, {"text": "b"}, {"text": "c", "correct": True}, {"text": "d"}]),
             Snapshot("short", "z", [], answer_text="-1,5")]
    assert sheet_problem(snaps, "12", lay) is None
    assert sheet_problem(snaps, "12") is not None          # Mẫu 40: only A–D
    key = grading_key(snaps, lay)
    assert key["q1"] == "E" and key["s1"] == "-1.5" and [key[f"t{i}"] for i in range(1, 5)] == ["Đ", "S", "Đ", "S"]
    assert sheet_problem(snaps * 3, "12", lay) is not None  # 3 trắc nghiệm on a 2-câu sheet
