"""
Trộn đề giữ nguyên định dạng Word (docx_mixer) + đáp án chọn trên web.

The mã đề must be the teacher's own paragraphs rearranged: pictures/formulas
move with their option, the answer marks and "Đáp án:" lines are gone, and
the options/answers agree with the text snapshots the answer key is built from.
"""
import base64
import io
import json
import zipfile

from docx import Document
from docx.shared import Pt, RGBColor

from app.services.docx_mixer import version_docx_keep_format
from app.services.rich_objects import TOKEN_RE
from app.services.exam_mixer import answer_key, build_versions
from app.services.question_io import parse_docx, parse_docx_layout

# 1×1 PNG
PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg==")


def _docx(marked: bool = True) -> bytes:
    doc = Document()
    doc.add_paragraph("ĐỀ KIỂM TRA TOÁN").runs[0].bold = True
    doc.add_paragraph("<g3>")
    for n in range(1, 6):
        doc.add_paragraph(f"Câu {n}. Câu hỏi số {n}?")
        p = doc.add_paragraph()
        for k, letter in enumerate("ABCD"):
            if k:
                p.add_run("\t")
            lab = p.add_run(f"{letter}. ")
            if marked and k == n % 4:
                lab.font.color.rgb = RGBColor(0xFF, 0, 0)
                lab.underline = True
            p.add_run(f"đáp án {letter}{n}")
            if letter == "C":          # a picture inside option C, like a MathType formula
                p.add_run().add_picture(io.BytesIO(PNG), width=Pt(8))
    doc.add_paragraph("Câu 6. Câu có dòng đáp án?")
    doc.add_paragraph("A. một\tB. hai\tC. ba\tD. bốn")
    doc.add_paragraph("Đáp án: " + ("B" if marked else ""))
    doc.add_paragraph("PHẦN IV. TRẢ LỜI NGẮN")
    doc.add_paragraph("Câu 1. Bao nhiêu?")
    doc.add_paragraph("Đáp án: 2,5")
    doc.add_paragraph("----- HẾT -----")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _mix(src: bytes, seed: int = 5):
    qs, _, _ = parse_docx_layout(src, allow_unanswered=True)
    snaps = build_versions(qs, ["101"], seed=seed)["101"]
    plan = [{"src": s.src, "qtype": s.qtype, "perm": s.perm} for s in snaps]
    return qs, snaps, version_docx_keep_format(src, "101", plan)


def _pictures(data: bytes) -> int:
    return len(Document(io.BytesIO(data)).element.body.xpath(".//w:drawing"))


def test_keep_format_moves_options_and_hides_answers():
    src = _docx()
    qs, snaps, out = _mix(src)
    assert [q.answer for q in qs if q.qtype == "mcq"] == [1, 2, 3, 0, 1, 1]
    doc = Document(io.BytesIO(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Mã đề: 101" in text and "ĐỀ KIỂM TRA TOÁN" in text and "HẾT" in text
    assert "Đáp án" not in text and "<g3>" not in text          # no answer lines, no group lines
    assert "PHẦN IV. TRẢ LỜI NGẮN" in text
    # nothing red or underlined left: the đề doesn't give the answer away
    for r in doc.element.body.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr"):
        assert not r.xpath("./w:u"), "underline left in the đề"
        assert all(c.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val") != "FF0000"
                   for c in r.xpath("./w:color"))
    assert _pictures(out) == _pictures(src) == 5
    # the đề students get = the snapshot the đáp án is computed from
    again, _ = parse_docx(out, allow_unanswered=True)
    mcq_out = [q for q in again if q.qtype == "mcq"]
    mcq_snap = [s for s in snaps if s.qtype == "mcq"]
    assert [[o["text"] for o in q.options] for q in mcq_out] == [[o["text"] for o in s.options] for s in mcq_snap]
    assert [q.content for q in mcq_out] == [s.content for s in mcq_snap]
    assert [p.text.split(".")[0] for p in doc.paragraphs if p.text.startswith("Câu")][:6] == \
        ["Câu 1", "Câu 2", "Câu 3", "Câu 4", "Câu 5", "Câu 6"]


def test_keep_format_picture_moves_with_its_option():
    src = _docx()
    _, snaps, out = _mix(src, seed=11)
    doc = Document(io.BytesIO(out))
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    for s in [s for s in snaps if s.qtype == "mcq" and s.content.startswith("Câu hỏi")]:
        words = lambda t: TOKEN_RE.sub("", t).strip()
        p = next(p for p in doc.paragraphs if "đáp án A" in p.text and words(s.options[0]["text"]) in p.text)
        # the picture is in the run right after the text of the option that was C
        pos_c = [o["text"].startswith("đáp án C") for o in s.options].index(True)
        runs = list(p._p.iterchildren(W + "r"))
        pic = next(i for i, r in enumerate(runs) if r.xpath(".//w:drawing"))
        before = "".join(r.xpath("string(.//w:t)") for r in runs[:pic])
        # the option text carries the picture as a "[[ct:…]]" token; compare the words
        assert "[[ct:" in s.options[pos_c]["text"]
        assert before.rstrip().endswith(TOKEN_RE.sub("", s.options[pos_c]["text"]).strip()), (before, s.options)


def test_unanswered_questions_are_kept_for_the_web():
    src = _docx(marked=False)
    qs, warnings = parse_docx(src)
    assert qs[-1].qtype == "short" and all(q.qtype != "mcq" for q in qs)   # default: skipped
    assert sum("chưa đánh dấu" in w for w in warnings) == 6
    qs, _ = parse_docx(src, allow_unanswered=True)
    assert [q.answer for q in qs if q.qtype == "mcq"] == [-1] * 6
    snaps = build_versions(qs, ["101"], seed=1)["101"]
    assert answer_key(snaps)["mcq"] == ["?"] * 6


def test_api_keep_format_and_picked_answers(client):
    h = client.headers_for(1)
    src = _docx(marked=False)
    f = lambda: {"file": ("toan.docx", src, "application/octet-stream")}

    info = client.post("/api/v1/exam-papers/parse-file", headers=h, files=f()).json()
    assert info["keeps_format"] is True
    assert [u["i"] for u in info["unanswered"]] == [0, 1, 2, 3, 4, 5]
    assert info["available"]["mcq"] == 0 and info["available_all"]["mcq"] == 6

    # chấm bằng phiếu: only the questions given an answer on the page are used
    r = client.post("/api/v1/exam-papers/from-file", headers=h, files=f(), data={
        "name": "Toán", "num_versions": "2", "answers_json": json.dumps({"0": 2, "3": 1, "5": 0})})
    assert r.status_code == 201, r.text
    p = r.json()
    assert p["counts"]["mcq"] == 3 and p["gradable"] is True
    assert p["notes"][0].startswith("Giữ nguyên định dạng Word")
    assert any("3 câu chưa có đáp án" in n for n in p["notes"])
    z = zipfile.ZipFile(io.BytesIO(client.get(f"/api/v1/exam-papers/{p['id']}/zip", headers=h).content))
    de = z.read(sorted(n for n in z.namelist() if n.startswith("Ma de"))[0])
    assert "Mã đề: " in "\n".join(x.text for x in Document(io.BytesIO(de)).paragraphs)
    assert _pictures(de) == 2           # questions 1 and 4 have a picture in option C; 6 has none

    # đề chỉ để in: every question, "?" where no answer was picked
    r = client.post("/api/v1/exam-papers/from-file", headers=h, files=f(), data={
        "name": "In", "for_sheet": "false", "answers_json": json.dumps({"1": 3})})
    p = r.json()
    assert p["counts"]["mcq"] == 6 and p["gradable"] is False
    assert sorted(p["versions"][0]["answer_key"]["mcq"]).count("?") == 5

    # several files → pooled as text, no format keeping
    r = client.post("/api/v1/exam-papers/parse-file", headers=h,
                    files=[("file", ("a.docx", src, "application/octet-stream")),
                           ("file", ("b.docx", src, "application/octet-stream"))])
    assert r.json()["keeps_format"] is False


def test_bank_import_picks_answers_on_the_page(client):
    h = client.headers_for(1)
    cid = client.post("/api/v1/question-bank/categories", headers=h, json={"name": "Toán"}).json()["id"]
    url = f"/api/v1/question-bank/categories/{cid}/import"
    f = lambda: {"file": ("toan.docx", _docx(marked=False), "application/octet-stream")}

    pre = client.post(url + "?dry_run=true", headers=h, files=f()).json()
    assert [u["i"] for u in pre["unanswered_list"]] == [0, 1, 2, 3, 4, 5]
    assert pre["unanswered_list"][0]["number"] == "Câu 1." and len(pre["unanswered_list"][0]["options"]) == 4
    assert [u["mcq_no"] for u in pre["unanswered_list"]] == [1, 2, 3, 4, 5, 6]
    assert [m["answer"] for m in pre["mcq_all"]] == [-1] * 6

    # two picked on the page now, the rest picked later
    r = client.post(url, headers=h, files=f(), data={"answers_json": json.dumps({"0": 2, "5": 1})}).json()
    assert r["new"] == 7 and r["unanswered"] == 4 and r["answered"] == 0

    # same file again: nothing new, but the bank's 4 unanswered still show and take an answer
    pre = client.post(url + "?dry_run=true", headers=h, files=f()).json()
    assert pre["new"] == 0 and [u["i"] for u in pre["unanswered_list"]] == [1, 2, 3, 4]
    r = client.post(url, headers=h, files=f(), data={"answers_json": json.dumps({"1": 0, "3": 3})}).json()
    assert r["new"] == 0 and r["answered"] == 2 and r["unanswered"] == 2

    qs = client.get(f"/api/v1/question-bank/categories/{cid}/questions", headers=h).json()
    mcq = sorted((q["content"], q["answer"]) for q in qs if q["qtype"] == "mcq")
    assert [a for _, a in mcq] == [1, 2, 0, -1, 3, -1]       # "Câu có dòng…" (câu 6) sorts first

    bad = client.post(url, headers=h, files=f(), data={"answers_json": "{oops"})
    assert bad.status_code == 422


def test_bank_keeps_pictures_and_formulas(client):
    """A picture/formula in a Word câu: kept as a "[[ct:…]]" token + asset in the
    ngân hàng, shown on the page (SVG), put back as the real object in the Word
    export and in mã đề mixed from the bank; Moodle .txt says "[công thức]"."""
    h = client.headers_for(1)
    src = _docx()
    qs, _ = parse_docx(src)
    c_opts = [o["text"] for q in qs if q.qtype == "mcq" for o in q.options if o["text"].startswith("đáp án C")]
    assert c_opts and all(TOKEN_RE.search(t) for t in c_opts)
    assert all(q.assets for q in qs if q.qtype == "mcq" and any(TOKEN_RE.search(o["text"]) for o in q.options))

    cid = client.post("/api/v1/question-bank/categories", headers=h, json={"name": "Có hình"}).json()["id"]
    r = client.post(f"/api/v1/question-bank/categories/{cid}/import", headers=h,
                    files={"file": ("hinh.docx", src, "application/octet-stream")}).json()
    assert r["new"] == len(qs) and not any("bị mất" in w for w in r["warnings"])

    asset_id = TOKEN_RE.search(c_opts[0]).group(1)
    svg = client.get(f"/api/v1/question-bank/assets/{asset_id}.svg")
    if r["formula_note"] is None:          # LibreOffice-free: pictures need none
        assert svg.status_code == 200 and svg.text.startswith("<svg")
    assert client.get("/api/v1/question-bank/assets/ffffffffffffffffffffffff.svg").status_code == 404

    word = client.get(f"/api/v1/question-bank/categories/{cid}/export?format=docx", headers=h).content
    assert _pictures(word) == 5 and "[công thức]" not in "\n".join(p.text for p in Document(io.BytesIO(word)).paragraphs)
    again, _ = parse_docx(word)
    assert [q.options for q in again if q.qtype == "mcq"] == [q.options for q in qs if q.qtype == "mcq"]

    txt = client.get(f"/api/v1/question-bank/categories/{cid}/export?format=txt", headers=h).text
    assert "[công thức]" in txt and "[[ct:" not in txt

    p = client.post("/api/v1/exam-papers/from-bank", headers=h, json={
        "name": "Từ ngân hàng", "category_ids": [cid], "counts": {"mcq": 6, "tf": 0, "short": 0},
        "num_versions": 1, "start_code": "101", "shuffle_questions": True, "shuffle_options": True,
        "exam_id": None, "for_sheet": False}).json()
    de = client.get(f"/api/v1/exam-papers/{p['id']}/versions/101/docx", headers=h)
    if de.status_code == 404:
        de = client.get(f"/api/v1/exam-papers/{p['id']}/zip", headers=h)
        z = zipfile.ZipFile(io.BytesIO(de.content))
        de_bytes = z.read(sorted(n for n in z.namelist() if n.startswith("Ma de"))[0])
    else:
        de_bytes = de.content
    assert _pictures(de_bytes) == 5


def test_superscript_and_floating_figure():
    """130⁰ stays 130⁰ (not "1300"); a floating (anchored) figure goes with
    the câu's content, not lost or stuck on an option."""
    from docx.oxml import parse_xml

    doc = Document()
    p = doc.add_paragraph("Câu 1. Số đo góc AOB là:")
    q = doc.add_paragraph()
    for k, letter in enumerate("ABCD"):
        q.add_run(("\t" if k else "") + f"{letter}. 1{3 + k}0")
        q.add_run("0").font.superscript = True
    q.runs[0].font.color.rgb = RGBColor(0xFF, 0, 0)
    q.runs[0].underline = True
    # a floating (anchored) picture on the câu line
    p2 = doc.add_paragraph()
    p2.add_run("Câu 2. Trong hình bên, x bằng:")
    run = p2.runs[0]
    run.add_picture(io.BytesIO(PNG), width=Pt(40))
    inline = run._r.xpath(".//wp:inline")[0]
    anchor = parse_xml(
        '<wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="1" behindDoc="0" locked="0" '
        'layoutInCell="1" allowOverlap="1"><wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="column"><wp:posOffset>0</wp:posOffset></wp:positionH>'
        '<wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV></wp:anchor>')
    for child in list(inline):
        anchor.append(child)
    inline.getparent().replace(inline, anchor)
    doc.add_paragraph("A. 1\tB. 2\tC. 3\tD. 4")
    doc.add_paragraph("Đáp án: B")
    buf = io.BytesIO()
    doc.save(buf)

    qs, warnings = parse_docx(buf.getvalue())
    assert [o["text"] for o in qs[0].options] == ["130⁰", "140⁰", "150⁰", "160⁰"] and qs[0].answer == 0
    assert TOKEN_RE.search(qs[1].content) and not any(TOKEN_RE.search(o["text"]) for o in qs[1].options)
    assert len(qs[1].assets) == 1 and not any("bị mất" in w for w in warnings)


def test_formulas_imported_without_libreoffice_are_drawn_later(client, monkeypatch):
    """Imported while the server had no LibreOffice: kept, the page is told;
    once it's there, opening the danh mục draws them."""
    import pytest
    from docx.oxml import parse_xml
    from app.services import rich_objects as ro

    real = ro.soffice_path()
    doc = Document()
    p = doc.add_paragraph("Câu 1. Giá trị của ")
    p._p.append(parse_xml(
        '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
        '<m:r><m:t>x+1=2</m:t></m:r></m:oMath>'))
    p.add_run(" là:")
    doc.add_paragraph("A. 1\tB. 2\tC. 3\tD. 4")
    doc.add_paragraph("Đáp án: A")
    buf = io.BytesIO()
    doc.save(buf)

    h = client.headers_for(1)
    cid = client.post("/api/v1/question-bank/categories", headers=h, json={"name": "Chưa có LibreOffice"}).json()["id"]
    monkeypatch.setattr(ro, "soffice_path", lambda: None)
    r = client.post(f"/api/v1/question-bank/categories/{cid}/import", headers=h,
                    files={"file": ("pt.docx", buf.getvalue(), "application/octet-stream")}).json()
    assert r["new"] == 1 and r["formula_note"]
    q = client.get(f"/api/v1/question-bank/categories/{cid}/questions", headers=h).json()[0]
    aid = TOKEN_RE.search(q["content"]).group(1)
    assert client.get(f"/api/v1/question-bank/assets/{aid}.svg").status_code == 404
    # the Word export still has the real equation
    word = client.get(f"/api/v1/question-bank/categories/{cid}/export?format=docx", headers=h).content
    assert len(Document(io.BytesIO(word)).element.body.xpath(".//m:oMath")) == 1

    if real is None:
        pytest.skip("no LibreOffice here to draw it")
    monkeypatch.setattr(ro, "soffice_path", lambda: real)
    client.get(f"/api/v1/question-bank/categories/{cid}/questions", headers=h)
    svg = client.get(f"/api/v1/question-bank/assets/{aid}.svg")
    assert svg.status_code == 200 and svg.text.startswith("<svg")
