"""
test_question_bank.py
=====================
Question bank: Aiken/Word parsers and the /question-bank API.
"""
import io
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent.parent))

from app.services.question_io import export_aiken, export_docx, parse_aiken, parse_docx

AIKEN_SAMPLE = """\
Do not touch any computer components with a magnetic screwdriver.
A) True
B) False
ANSWER: A
Screwdriver with magnetic tip
ANSWER:
Which is the fastest cache in a computer?
A) L1
B) L2
C) L3
D) L4
ANSWER: A
What size power supply is needed?
A) 500 watts
B) 200 watts
C) 300 watts
D) 400 watts
E) 600 watts
ANSWER: E
"""


# ── Parsers ──────────────────────────────────────────────────────────────────

def test_parse_aiken_moodle_sample():
    qs, warnings = parse_aiken(AIKEN_SAMPLE)
    assert [q.content for q in qs] == [
        "Do not touch any computer components with a magnetic screwdriver.",
        "Which is the fastest cache in a computer?",
        "What size power supply is needed?",
    ]
    assert [q.answer for q in qs] == [0, 0, 4]
    assert qs[0].shuffle_options is False       # True/False keeps its order
    assert qs[1].shuffle_options is True
    assert len(warnings) == 1 and "Screwdriver with magnetic tip" in warnings[0]


def test_parse_aiken_multiline_question_and_bad_answer():
    qs, warnings = parse_aiken("Line one\nline two\nA. x\nB. y\nANSWER: B\n\nQ2\nA) x\nB) y\nANSWER: C\n")
    assert len(qs) == 1 and qs[0].content == "Line one\nline two" and qs[0].answer == 1
    assert len(warnings) == 1 and "ANSWER không hợp lệ" in warnings[0]


def test_aiken_roundtrip():
    qs, _ = parse_aiken(AIKEN_SAMPLE)
    again, warnings = parse_aiken(export_aiken(qs))
    assert warnings == [] and again == qs


def test_docx_roundtrip():
    qs, _ = parse_aiken(AIKEN_SAMPLE)
    qs[1].options[3]["fixed"] = True
    again, warnings = parse_docx(export_docx(qs, title="Đề"))
    assert warnings == []
    assert [(q.content, q.options, q.answer) for q in again] == [(q.content, q.options, q.answer) for q in qs]


def test_parse_docx_youngmix_rules():
    from docx import Document
    from docx.shared import RGBColor

    doc = Document()
    doc.add_paragraph("ĐỀ KIỂM TRA")            # ignored: before the first question
    doc.add_paragraph("<g1>")
    doc.add_paragraph("Câu 1: 2 + 2 = ?")
    p = doc.add_paragraph()
    p.add_run("A. 3\t")
    p.add_run("B.").underline = True
    p.add_run(" 4\tC. 5\t#D. Không có đáp án")
    doc.add_paragraph("<g3>")
    doc.add_paragraph("Question 2. Which vitamin?")
    doc.add_paragraph("A. vitamin A. is not it")
    doc.add_paragraph().add_run("B. C").font.color.rgb = RGBColor(0xFF, 0, 0)
    doc.add_paragraph("Câu 3. Unmarked")
    doc.add_paragraph("A. x")
    doc.add_paragraph("B. y")
    buf = io.BytesIO()
    doc.save(buf)

    qs, warnings = parse_docx(buf.getvalue())
    assert len(qs) == 2
    assert qs[0].content == "2 + 2 = ?"
    assert [o["text"] for o in qs[0].options] == ["3", "4", "5", "Không có đáp án"]
    assert qs[0].options[3]["fixed"] is True and qs[0].answer == 1
    assert qs[0].shuffle_options is False       # <g1>: options keep their order
    assert [o["text"] for o in qs[1].options] == ["vitamin A. is not it", "C"]
    assert qs[1].answer == 1 and qs[1].shuffle_options is True
    assert len(warnings) == 1 and "Phần I-II câu 3" in warnings[0]


def _docx_three_parts() -> bytes:
    from docx import Document
    from docx.shared import RGBColor

    doc = Document()
    doc.add_paragraph("PHẦN I. TRẮC NGHIỆM")
    doc.add_paragraph("Câu 1. Bộ nhớ đệm nào nhanh nhất?")
    doc.add_paragraph("A. L1\tB. L2\tC. L3\tD. L4")
    doc.add_paragraph("Đáp án: A")                       # key line instead of colouring
    doc.add_paragraph("PHẦN II. ĐÚNG/SAI")
    doc.add_paragraph("Câu 1. Cho các phát biểu về CPU:")
    doc.add_paragraph().add_run("a) CPU là bộ xử lý trung tâm").font.color.rgb = RGBColor(0xFF, 0, 0)
    doc.add_paragraph("b) RAM lưu dữ liệu vĩnh viễn")
    doc.add_paragraph().add_run("c) L1 nhanh hơn L2").underline = True
    doc.add_paragraph("d) SSD là bộ nhớ quang")
    doc.add_paragraph("Câu 2. Phát biểu về mạng:")
    for t in ("a) x", "b) y", "c) z", "d) w"):
        doc.add_paragraph(t)
    doc.add_paragraph("Đáp án: S Đ Đ S")
    doc.add_paragraph("Câu 3. Chỉ có 3 ý:")
    for t in ("a) x", "b) y", "c) z"):
        doc.add_paragraph(t)
    doc.add_paragraph("PHẦN III. TRẢ LỜI NGẮN")
    doc.add_paragraph("Câu 1. Ổ 500 GB dùng 125 GB, còn trống bao nhiêu %?")
    doc.add_paragraph("Đáp án: 75")
    doc.add_paragraph("Câu 2. Giá trị của -3/2?")
    doc.add_paragraph("Đáp án: -1.5")                    # "." accepted → "-1,5"
    doc.add_paragraph("Câu 3. Quá dài")
    doc.add_paragraph("Đáp án: 12345")
    doc.add_paragraph("Câu 4. Thiếu đáp án")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_parse_docx_three_parts():
    qs, warnings = parse_docx(_docx_three_parts())
    assert [q.qtype for q in qs] == ["mcq", "tf", "tf", "short", "short"]
    mcq, tf1, tf2, s1, s2 = qs
    assert mcq.answer == 0 and [o["text"] for o in mcq.options] == ["L1", "L2", "L3", "L4"]
    assert [o["correct"] for o in tf1.options] == [True, False, True, False]
    assert tf1.options[1]["text"] == "RAM lưu dữ liệu vĩnh viễn" and tf1.shuffle_options is False
    assert [o["correct"] for o in tf2.options] == [False, True, True, False]
    assert (s1.answer_text, s2.answer_text) == ("75", "-1,5")
    assert s1.content == "Ổ 500 GB dùng 125 GB, còn trống bao nhiêu %?"
    assert len(warnings) == 3
    assert "Phần III câu 3" in warnings[0] and "4 ý" in warnings[0]
    assert "Phần IV câu 3" in warnings[1] and "12345" in warnings[1]
    assert "Phần IV câu 4" in warnings[2] and "thiếu" in warnings[2]


def _docx_headings(headings: tuple[str, str, str]) -> bytes:
    from docx import Document
    doc = Document()
    doc.add_paragraph(headings[0])
    doc.add_paragraph("Câu 1. TN?")
    doc.add_paragraph("A. x\tB. y\tC. z\tD. w")
    doc.add_paragraph("Đáp án: C")
    doc.add_paragraph(headings[1])
    doc.add_paragraph("Câu 1. Đúng sai?")
    for t in ("a) x", "b) y", "c) z", "d) w"):
        doc.add_paragraph(t)
    doc.add_paragraph("Đáp án: Đ S S Đ")
    doc.add_paragraph(headings[2])
    doc.add_paragraph("Câu 1. Số?")
    doc.add_paragraph("Đáp án: 2,5")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.mark.parametrize("headings", [
    ("PHẦN I-II. TRẮC NGHIỆM", "PHẦN III. ĐÚNG/SAI", "PHẦN IV. TRẢ LỜI NGẮN"),   # VJU sheet naming
    ("PHẦN I-II", "PHẦN III", "PHẦN IV"),                                     # bare, VJU numbers
    ("PHẦN I. TRẮC NGHIỆM", "PHẦN II. ĐÚNG/SAI", "PHẦN III. TRẢ LỜI NGẮN"),   # Bộ GD naming
    ("PHẦN I", "PHẦN II", "PHẦN III"),                                        # bare, Bộ GD numbers
    ("Phần 1", "Phần 2", "Phần 3"),
])
def test_part_headings_either_numbering(headings):
    """Đề files numbered like the VJU sheet (I-II / III / IV) or the Bộ GD way
    (I / II / III) both read as trắc nghiệm, Đúng/Sai, trả lời ngắn."""
    qs, warnings = parse_docx(_docx_headings(headings))
    assert warnings == []
    assert [q.qtype for q in qs] == ["mcq", "tf", "short"]
    assert qs[0].answer == 2
    assert [o["correct"] for o in qs[1].options] == [True, False, False, True]
    assert qs[2].answer_text == "2,5"


def test_docx_three_parts_roundtrip():
    qs, _ = parse_docx(_docx_three_parts())
    again, warnings = parse_docx(export_docx(qs, title="Đề"))
    assert warnings == [] and again == qs


def test_normalize_short_answer():
    from app.services.question_io import normalize_short_answer as n
    assert [n(x) for x in ("75", "-1.5", " 0,25 ", "2025", "-12")] == ["75", "-1,5", "0,25", "2025", "-12"]
    assert [n(x) for x in ("12345", "-1,25", "1,2,3", "abc", "", "1-2", ",5")] == [None] * 7


def test_aiken_export_skips_non_mcq():
    qs, _ = parse_docx(_docx_three_parts())
    text = export_aiken(qs)
    assert text.count("ANSWER:") == 1 and "CPU" not in text


def test_parse_docx_rejects_non_docx():
    qs, warnings = parse_docx(b"not a docx")
    assert qs == [] and ".docx" in warnings[0]


# ── API ──────────────────────────────────────────────────────────────────────


def test_api_crud_import_export(client):
    h1, h2 = client.headers_for(1), client.headers_for(2)

    r = client.post("/api/v1/question-bank/categories", json={"name": "Tin học"}, headers=h1)
    assert r.status_code == 201
    cid = r.json()["id"]

    # Other teachers can't see it
    assert client.get("/api/v1/question-bank/categories", headers=h2).json() == []
    assert client.get(f"/api/v1/question-bank/categories/{cid}/questions", headers=h2).status_code == 404

    # Dry-run import saves nothing
    files = {"file": ("q.txt", AIKEN_SAMPLE.encode(), "text/plain")}
    r = client.post(f"/api/v1/question-bank/categories/{cid}/import?dry_run=true", files=files, headers=h1)
    assert r.json() == {"dry_run": True, "found": 3, "new": 3, "duplicates": 0, "unanswered": 0,
                        "answered": 0, "unanswered_list": [], "mcq_all": r.json()["mcq_all"], "duplicate_list": [],
                        "formula_note": None,
                        "warnings": r.json()["warnings"]}
    assert client.get(f"/api/v1/question-bank/categories/{cid}/questions", headers=h1).json() == []

    r = client.post(f"/api/v1/question-bank/categories/{cid}/import", files=files, headers=h1)
    assert r.json()["new"] == 3
    # Re-importing the same file: all duplicates
    r = client.post(f"/api/v1/question-bank/categories/{cid}/import", files=files, headers=h1)
    assert r.json()["new"] == 0 and r.json()["duplicates"] == 3
    assert all(d.endswith("đã có trong danh mục") for d in r.json()["duplicate_list"])
    # the same câu twice in one file: the preview names both
    twice = (AIKEN_SAMPLE + "\n".join(AIKEN_SAMPLE.splitlines()[:4]) + "\n").encode()
    cid2 = client.post("/api/v1/question-bank/categories", json={"name": "Trùng"}, headers=h1).json()["id"]
    r = client.post(f"/api/v1/question-bank/categories/{cid2}/import?dry_run=true",
                    files={"file": ("q.txt", twice, "text/plain")}, headers=h1).json()
    assert r["duplicates"] == 1 and "trùng Câu 1 (dòng 1) trong cùng file" in r["duplicate_list"][0], r
    client.delete(f"/api/v1/question-bank/categories/{cid2}", headers=h1)

    qs = client.get(f"/api/v1/question-bank/categories/{cid}/questions", headers=h1).json()
    assert len(qs) == 3
    assert client.get("/api/v1/question-bank/categories", headers=h1).json()[0]["question_count"] == 3
    assert len(client.get(f"/api/v1/question-bank/categories/{cid}/questions?q=cache", headers=h1).json()) == 1

    # Manual add / edit / validation / delete
    body = {"content": "1 + 1 = ?", "options": [{"text": "1"}, {"text": "2"}], "answer": 1}
    r = client.post(f"/api/v1/question-bank/categories/{cid}/questions", json=body, headers=h1)
    assert r.status_code == 201
    qid = r.json()["id"]
    body["answer"] = 5
    assert client.put(f"/api/v1/question-bank/questions/{qid}", json=body, headers=h1).status_code == 422
    body["answer"] = 0
    assert client.put(f"/api/v1/question-bank/questions/{qid}", json=body, headers=h1).json()["answer"] == 0
    assert client.delete(f"/api/v1/question-bank/questions/{qid}", headers=h2).status_code == 404
    assert client.delete(f"/api/v1/question-bank/questions/{qid}", headers=h1).status_code == 204

    # Export both formats
    r = client.get(f"/api/v1/question-bank/categories/{cid}/export?format=txt", headers=h1)
    assert r.status_code == 200 and "ANSWER: E" in r.text
    r = client.get(f"/api/v1/question-bank/categories/{cid}/export?format=docx", headers=h1)
    assert r.status_code == 200 and len(parse_docx(r.content)[0]) == 3

    # Unsupported file type
    r = client.post(f"/api/v1/question-bank/categories/{cid}/import",
                    files={"file": ("q.pdf", b"%PDF", "application/pdf")}, headers=h1)
    assert r.status_code == 400

    # Deleting the category deletes its questions
    assert client.delete(f"/api/v1/question-bank/categories/{cid}", headers=h1).status_code == 204
    assert client.get("/api/v1/question-bank/categories", headers=h1).json() == []


def test_api_sharing(client):
    h1, h2 = client.headers_for(1), client.headers_for(2)
    base = "/api/v1/question-bank"
    cid = client.post(f"{base}/categories", json={"name": "Toán"}, headers=h1).json()["id"]
    body = {"content": "1 + 1 = ?", "options": [{"text": "1"}, {"text": "2"}], "answer": 1}
    qid = client.post(f"{base}/categories/{cid}/questions", json=body, headers=h1).json()["id"]

    # Unknown account / sharing with the owner
    assert client.post(f"{base}/categories/{cid}/shares", json={"email": "nobody@x.vn"}, headers=h1).status_code == 404
    assert client.post(f"{base}/categories/{cid}/shares", json={"email": "t1@vju.ac.vn"}, headers=h1).status_code == 400

    # Share read-only (email case-insensitive)
    r = client.post(f"{base}/categories/{cid}/shares", json={"email": "T2@VJU.ac.vn"}, headers=h1)
    assert [(s["email"], s["permission"]) for s in r.json()] == [("t2@vju.ac.vn", "view")]
    cats = client.get(f"{base}/categories", headers=h2).json()
    assert [(c["id"], c["access"], c["owner_name"]) for c in cats] == [(cid, "view", "t1@vju.ac.vn")]
    assert len(client.get(f"{base}/categories/{cid}/questions", headers=h2).json()) == 1
    assert client.get(f"{base}/categories/{cid}/export?format=txt", headers=h2).status_code == 200
    # …but can't change anything
    assert client.post(f"{base}/categories/{cid}/questions", json=body, headers=h2).status_code == 403
    assert client.put(f"{base}/questions/{qid}", json=body, headers=h2).status_code == 403
    assert client.delete(f"{base}/questions/{qid}", headers=h2).status_code == 403
    files = {"file": ("q.txt", AIKEN_SAMPLE.encode(), "text/plain")}
    assert client.post(f"{base}/categories/{cid}/import", files=files, headers=h2).status_code == 403
    assert client.put(f"{base}/categories/{cid}", json={"name": "x"}, headers=h2).status_code == 403
    assert client.post(f"{base}/categories/{cid}/shares", json={"email": "t1@vju.ac.vn"}, headers=h2).status_code == 403

    # Copy into own bank → fully editable, independent of the original
    copy = client.post(f"{base}/categories/{cid}/copy", headers=h2).json()
    assert copy["access"] == "owner" and copy["owner_id"] == 2 and copy["question_count"] == 1
    cq = client.get(f"{base}/categories/{copy['id']}/questions", headers=h2).json()[0]
    assert client.delete(f"{base}/questions/{cq['id']}", headers=h2).status_code == 204
    assert len(client.get(f"{base}/categories/{cid}/questions", headers=h1).json()) == 1

    # Upgrade to edit (sharing again updates the permission, no duplicate row)
    r = client.post(f"{base}/categories/{cid}/shares", json={"email": "t2@vju.ac.vn", "permission": "edit"}, headers=h1)
    assert [(s["permission"]) for s in r.json()] == ["edit"]
    assert client.put(f"{base}/questions/{qid}", json=body, headers=h2).status_code == 200
    assert client.delete(f"{base}/categories/{cid}", headers=h2).status_code == 403   # still owner-only

    # Revoke
    share_id = r.json()[0]["id"]
    assert client.delete(f"{base}/categories/{cid}/shares/{share_id}", headers=h1).status_code == 204
    assert client.get(f"{base}/categories/{cid}/questions", headers=h2).status_code == 404
    assert [c["id"] for c in client.get(f"{base}/categories", headers=h2).json()] == [copy["id"]]

    # Deleting the category also removes its shares
    client.post(f"{base}/categories/{cid}/shares", json={"email": "t2@vju.ac.vn"}, headers=h1)
    assert client.delete(f"{base}/categories/{cid}", headers=h1).status_code == 204
    assert [c["id"] for c in client.get(f"{base}/categories", headers=h2).json()] == [copy["id"]]


def test_api_user_state_survives_logout(client):
    """Answer keys saved on the server are per-account and survive logout."""
    h1, h2 = client.headers_for(1), client.headers_for(2)
    key = {"answers": {"cau1": "A"}, "scoring": {"correct": 1, "wrong": 0, "blank": 0}, "updatedAt": "x"}

    assert client.get("/api/v1/me/state", headers=h1).json() == {}
    assert client.put("/api/v1/me/state/vju_answer_key", json=key, headers=h1).status_code == 204
    assert client.put("/api/v1/me/state/vju_answer_key_library", json=[{"id": "1"}], headers=h1).status_code == 204
    # Overwrite keeps one row
    key["answers"]["cau2"] = "B"
    client.put("/api/v1/me/state/vju_answer_key", json=key, headers=h1)
    assert client.get("/api/v1/me/state", headers=h1).json() == {
        "vju_answer_key": key, "vju_answer_key_library": [{"id": "1"}]}

    assert client.get("/api/v1/me/state", headers=h2).json() == {}       # other account sees nothing
    assert client.put("/api/v1/me/state/whatever", json={}, headers=h1).status_code == 404
    assert client.get("/api/v1/me/state").status_code in (401, 403)

    # null deletes
    client.put("/api/v1/me/state/vju_answer_key_library", content="null",
               headers={**h1, "Content-Type": "application/json"})
    assert list(client.get("/api/v1/me/state", headers=h1).json()) == ["vju_answer_key"]


def test_shared_template_seeded_and_pinned_id(client, tmp_path, monkeypatch):
    """Fresh DB → Mẫu 40 installed automatically and its id served by
    /custom-forms/pinned (was hard-coded to production's id 2 in the frontend)."""
    from sqlalchemy.orm import sessionmaker
    from app.database import get_db
    from app.models.template import Template
    from app.services import shared_templates as st

    monkeypatch.setattr(st, "MAU40_DEST_TPL", tmp_path / "shared_40tn_dungsai.template.json")
    monkeypatch.setattr(st, "MAU40_DEST_AREAS", tmp_path / "shared_40tn_dungsai.areas.json")
    monkeypatch.setattr(st, "DEST_DIR", tmp_path)

    db = next(client.app.dependency_overrides[get_db]())
    # An unrelated custom template takes id 1, so Mẫu 40 gets a different id than prod's
    db.add(Template(name="Khác", type="custom", version="1.0", file_path="x.json", is_default=False))
    db.commit()
    assert client.get("/api/v1/custom-forms/pinned", headers=client.headers_for(1)).json() == {"mau40": None}

    st.ensure_shared_templates(db)
    tpl = st.find_mau40(db)
    assert tpl is not None and tpl.is_default and (tmp_path / "shared_40tn_dungsai.template.json").exists()
    assert client.get("/api/v1/custom-forms/pinned", headers=client.headers_for(2)).json() == {"mau40": tpl.id}
    assert client.get(f"/api/v1/custom-forms/{tpl.id}", headers=client.headers_for(2)).status_code == 200

    # Running again never duplicates or touches the existing row
    tpl.name = "Đổi tên bởi admin"
    db.commit()
    st.ensure_shared_templates(db)
    assert db.query(Template).filter(Template.type == "custom").count() == 2
    assert st.find_mau40(db).name == "Đổi tên bởi admin"   # still found by file name after a rename


def test_api_three_question_types(client):
    h = client.headers_for(1)
    base = "/api/v1/question-bank"
    cid = client.post(f"{base}/categories", json={"name": "Đề 2025"}, headers=h).json()["id"]
    post = lambda body: client.post(f"{base}/categories/{cid}/questions", json=body, headers=h)

    r = post({"qtype": "mcq", "content": "1+1?", "options": [{"text": "1"}, {"text": "2"}], "answer": 1})
    assert r.status_code == 201 and r.json()["options"] == [{"text": "1", "fixed": False, "correct": None},
                                                           {"text": "2", "fixed": False, "correct": None}]
    tf = {"qtype": "tf", "content": "CPU:", "options": [
        {"text": "a", "correct": True}, {"text": "b"}, {"text": "c", "correct": True}, {"text": "d"}]}
    r = post(tf)
    assert r.status_code == 201
    assert [o["correct"] for o in r.json()["options"]] == [True, False, True, False]
    assert r.json()["shuffle_options"] is False
    r = post({"qtype": "short", "content": "Còn trống bao nhiêu %?", "answer_text": "-1.5", "options": [{"text": "x"}]})
    assert r.status_code == 201 and r.json()["answer_text"] == "-1,5" and r.json()["options"] == []

    # Validation per type
    assert post({**tf, "options": tf["options"][:3]}).status_code == 422
    assert post({"qtype": "short", "content": "x", "answer_text": "12345"}).status_code == 422
    assert post({"qtype": "short", "content": "x"}).status_code == 422

    cat = client.get(f"{base}/categories", headers=h).json()[0]
    assert cat["question_count"] == 3 and cat["type_counts"] == {"mcq": 1, "tf": 1, "short": 1}
    assert [q["qtype"] for q in client.get(f"{base}/categories/{cid}/questions?qtype=tf", headers=h).json()] == ["tf"]

    # Word export → import into a new category keeps all 3 types; .txt keeps only mcq
    r = client.get(f"{base}/categories/{cid}/export?format=docx", headers=h)
    cid2 = client.post(f"{base}/categories", json={"name": "Nhập lại"}, headers=h).json()["id"]
    r = client.post(f"{base}/categories/{cid2}/import", headers=h,
                    files={"file": ("de.docx", r.content, "application/octet-stream")})
    assert r.json()["new"] == 3 and r.json()["warnings"] == []
    txt = client.get(f"{base}/categories/{cid}/export?format=txt", headers=h).text
    assert txt.count("ANSWER:") == 1

    # Copy keeps the types
    copy = client.post(f"{base}/categories/{cid}/copy", headers=h).json()
    assert copy["type_counts"] == {"mcq": 1, "tf": 1, "short": 1}


def test_migration_adds_question_type_columns(tmp_path, monkeypatch):
    """A `questions` table created before qtype/answer_text existed (local dev
    DBs) gets the columns on startup; old rows become "mcq"."""
    import sqlite3
    from sqlalchemy import create_engine
    import app.database as database

    path = tmp_path / "old.db"
    con = sqlite3.connect(path)
    con.execute("CREATE TABLE questions (id INTEGER PRIMARY KEY, category_id INTEGER, owner_id INTEGER, "
                "content TEXT, options_json TEXT, answer INTEGER, shuffle_options BOOLEAN, "
                "created_at DATETIME, updated_at DATETIME)")
    con.execute("INSERT INTO questions (id, category_id, owner_id, content, options_json, answer, shuffle_options) "
                "VALUES (1, 1, 1, 'cũ', '[]', 0, 1)")
    con.commit(); con.close()

    monkeypatch.setattr(database, "engine", create_engine(f"sqlite:///{path}"))
    database._migrate_question_columns()
    database._migrate_question_columns()   # idempotent

    con = sqlite3.connect(path)
    cols = [r[1] for r in con.execute("PRAGMA table_info(questions)")]
    assert "qtype" in cols and "answer_text" in cols
    assert con.execute("SELECT qtype, answer_text FROM questions").fetchall() == [("mcq", None)]


def test_parse_docx_plain_numbered_questions_and_formula_warning():
    """Many teachers' files number questions "1." instead of "Câu 1."; equation
    objects (MathType / Word equations) aren't text and must be flagged."""
    from docx import Document
    from docx.oxml import OxmlElement

    doc = Document()
    doc.add_paragraph("I/ CĂN THỨC")                 # section heading, ignored
    doc.add_paragraph("1. Căn bậc hai của 16 là:")
    doc.add_paragraph("\tA. 2\t\tB. 4\t\tC. 8\t\tD. 16")
    doc.add_paragraph("Đáp án: B")
    p = doc.add_paragraph("2) Tính giá trị biểu thức ")
    p._p.append(OxmlElement("w:object"))              # stands in for a MathType formula
    doc.add_paragraph("A. 1\tB. 2\tC. 3\tD. 4")
    doc.add_paragraph("Đáp án: C")
    buf = io.BytesIO()
    doc.save(buf)

    qs, warnings = parse_docx(buf.getvalue())
    assert [q.content for q in qs] == ["Căn bậc hai của 16 là:", "Tính giá trị biểu thức"]
    assert [q.answer for q in qs] == [1, 2]
    assert len(warnings) == 1 and "câu 2" in warnings[0] and "công thức" in warnings[0]


def test_import_without_answers_then_pick_on_the_page(client):
    """A practice file with no answer marked: imported anyway, listed as
    "chưa có đáp án", answers picked one by one; bank mixing for the sheet
    only uses the answered ones, a đề chỉ để in uses all ("?" in the key)."""
    from docx import Document
    doc = Document()
    for n in range(1, 5):
        doc.add_paragraph(f"{n}. Câu số {n}?")
        doc.add_paragraph(f"A. a{n}\tB. b{n}\tC. c{n}\tD. d{n}")
    buf = io.BytesIO()
    doc.save(buf)
    h = client.headers_for(1)
    cid = client.post("/api/v1/question-bank/categories", headers=h, json={"name": "Chưa đáp án"}).json()["id"]
    r = client.post(f"/api/v1/question-bank/categories/{cid}/import", headers=h,
                    files={"file": ("bt.docx", buf.getvalue(), "application/octet-stream")})
    assert r.status_code == 200, r.text
    assert (r.json()["new"], r.json()["unanswered"]) == (4, 4)
    cat = next(c for c in client.get("/api/v1/question-bank/categories", headers=h).json() if c["id"] == cid)
    assert cat["unanswered"] == 4
    todo = client.get(f"/api/v1/question-bank/categories/{cid}/questions?unanswered=true", headers=h).json()
    assert [q["answer"] for q in todo] == [-1] * 4

    for q, a in zip(todo[:2], (2, 0)):
        r = client.put(f"/api/v1/question-bank/questions/{q['id']}/answer", headers=h, json={"answer": a})
        assert r.status_code == 200 and r.json()["answer"] == a
    assert client.put(f"/api/v1/question-bank/questions/{todo[2]['id']}/answer", headers=h,
                      json={"answer": 7}).status_code == 422
    assert len(client.get(f"/api/v1/question-bank/categories/{cid}/questions?unanswered=true",
                          headers=h).json()) == 2
    # Aiken export needs an ANSWER line: only the 2 answered questions
    txt = client.get(f"/api/v1/question-bank/categories/{cid}/export?format=txt", headers=h).text
    assert txt.count("ANSWER:") == 2

    r = client.post("/api/v1/exam-papers/from-bank", headers=h, json={
        "name": "Phiếu", "category_ids": [cid], "counts": {"mcq": 2}, "num_versions": 1})
    assert r.status_code == 201 and any("2 câu chưa có đáp án" in n for n in r.json()["notes"])
    assert client.post("/api/v1/exam-papers/from-bank", headers=h, json={
        "name": "Phiếu", "category_ids": [cid], "counts": {"mcq": 3}, "num_versions": 1}).status_code == 422
    r = client.post("/api/v1/exam-papers/from-bank", headers=h, json={
        "name": "In", "category_ids": [cid], "counts": {"mcq": 4}, "num_versions": 1, "for_sheet": False})
    assert r.status_code == 201 and r.json()["versions"][0]["answer_key"]["mcq"].count("?") == 2
