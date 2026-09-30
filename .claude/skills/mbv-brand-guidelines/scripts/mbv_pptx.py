"""Dựng slide 16:9 chuẩn MBV bằng python-pptx.

    import sys; sys.path.insert(0, "<skill>/scripts")
    import mbv_pptx as S
    prs = S.new_deck()
    S.add_cover(prs, "Báo cáo quý III", "Khối Bán lẻ", "30/09/2026")
    S.add_content(prs, "Tổng quan", ["Ý 1", ("Ý con", 1), "Ý 2"])
    S.add_kpi(prs, "Chỉ số chính", [("1.250 tỷ", "Huy động"), ("18%", "Tăng trưởng")])
    S.add_table(prs, "Bảng số liệu", ["Chỉ tiêu", "Q2", "Q3"], [["A", "1", "2"]])
    S.add_closing(prs)
    prs.save("out.pptx")
"""
import copy
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import mbv_brand as B
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

SW, SH = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.7)
FOOTER_TEXT = "MBV · Ngân hàng TNHH MTV Việt Nam Hiện Đại"


def _rgb(h):
    return RGBColor.from_string(h.upper())


def new_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = SW, SH
    return prs


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _rect(slide, x, y, w, h, color):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = _rgb(color)
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def _style_run(run, size, bold=False, color=B.INK, italic=False):
    f = run.font
    f.name = B.FONT
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = _rgb(color)
    rpr = run._r.get_or_add_rPr()  # đảm bảo cả ký tự Á/Đông dùng cùng font
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", B.FONT)


def _text(slide, x, y, w, h, text, size, bold=False, color=B.INK,
          align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        _style_run(p.add_run(), size, bold, color, italic)
        p.runs[0].text = line
    return tb


def _picture(slide, path, x, y, width):
    """Chèn ảnh theo chiều rộng, tự giữ tỷ lệ (không bao giờ kéo méo)."""
    return slide.shapes.add_picture(path, x, y, width=width)


def _bullet(p, level, color=B.RED):
    ppr = p._p.get_or_add_pPr()
    indent = Emu(285750)
    ppr.set("marL", str(int(indent * (level + 1))))
    ppr.set("indent", str(-int(indent)))
    for tag, attrs in (("a:buClr", None), ("a:buFont", {"typeface": B.FONT}),
                       ("a:buChar", {"char": "•" if level == 0 else "–"})):
        el = etree.SubElement(ppr, qn(tag))
        if tag == "a:buClr":
            etree.SubElement(el, qn("a:srgbClr")).set("val", color)
        else:
            for k, v in attrs.items():
                el.set(k, v)


def _chrome(prs, slide, title):
    """Thanh đỏ trên cùng, tiêu đề, logo nhỏ, chân trang."""
    _rect(slide, 0, 0, SW, Inches(0.14), B.RED)
    _text(slide, MARGIN, Inches(0.5), Inches(10.4), Inches(0.9), title, 30, True, B.INK,
          anchor=MSO_ANCHOR.MIDDLE)
    _picture(slide, B.logo_path("color"), SW - MARGIN - Inches(1.3), Inches(0.62), Inches(1.3))
    _rect(slide, MARGIN, Inches(1.5), Inches(0.9), Inches(0.05), B.GOLD)
    _text(slide, MARGIN, SH - Inches(0.5), Inches(8), Inches(0.3), FOOTER_TEXT, 9, False, B.GREY)
    _text(slide, SW - MARGIN - Inches(1), SH - Inches(0.5), Inches(1), Inches(0.3),
          str(len(prs.slides)), 9, False, B.GREY, PP_ALIGN.RIGHT)


def add_cover(prs, title, subtitle="", meta=""):
    s = _blank(prs)
    _rect(s, 0, 0, SW, SH, B.RED)
    _rect(s, Inches(8.3), 0, SW - Inches(8.3), SH, B.DEEP_RED)
    _picture(s, B.star_path("color"), Inches(8.2), Inches(1.2), Inches(5.6))
    _picture(s, B.logo_path("white"), MARGIN, Inches(0.7), Inches(2.4))
    _text(s, MARGIN, Inches(2.9), Inches(7.4), Inches(2.0), title, 40, True, B.WHITE,
          anchor=MSO_ANCHOR.BOTTOM)
    _rect(s, MARGIN, Inches(5.1), Inches(1.2), Inches(0.06), B.GOLD)
    if subtitle:
        _text(s, MARGIN, Inches(5.3), Inches(7.4), Inches(0.6), subtitle, 18, False, B.WHITE)
    if meta:
        _text(s, MARGIN, SH - Inches(0.9), Inches(7), Inches(0.4), meta, 12, True, B.GOLD)
    return s


def add_section(prs, title, subtitle=""):
    s = _blank(prs)
    _rect(s, 0, 0, SW, SH, B.RED)
    _picture(s, B.star_path("color"), SW - Inches(4.2), SH - Inches(4.2), Inches(5.0))
    _text(s, MARGIN, Inches(2.6), Inches(9), Inches(1.6), title, 38, True, B.WHITE,
          anchor=MSO_ANCHOR.BOTTOM)
    _rect(s, MARGIN, Inches(4.4), Inches(1.2), Inches(0.06), B.GOLD)
    if subtitle:
        _text(s, MARGIN, Inches(4.6), Inches(8), Inches(0.8), subtitle, 18, False, B.WHITE)
    return s


def add_content(prs, title, bullets, size=20):
    """bullets: list[str | (str, level)]"""
    s = _blank(prs)
    _chrome(prs, s, title)
    tb = s.shapes.add_textbox(MARGIN, Inches(1.9), SW - 2 * MARGIN, SH - Inches(2.9))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, b in enumerate(bullets):
        text, level = (b, 0) if isinstance(b, str) else b
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        p.line_spacing = 1.2
        _style_run(p.add_run(), size - 3 * level, False, B.INK if level == 0 else B.GREY)
        p.runs[0].text = text
        _bullet(p, level)
    return s


def add_two_column(prs, title, left_title, left_bullets, right_title, right_bullets):
    s = _blank(prs)
    _chrome(prs, s, title)
    colw = (SW - 2 * MARGIN - Inches(0.6)) / 2
    for i, (ct, items) in enumerate(((left_title, left_bullets), (right_title, right_bullets))):
        x = MARGIN + i * (colw + Inches(0.6))
        _rect(s, x, Inches(1.9), colw, Inches(0.55), B.RED)
        _text(s, x + Inches(0.2), Inches(1.9), colw - Inches(0.4), Inches(0.55), ct, 16, True,
              B.WHITE, anchor=MSO_ANCHOR.MIDDLE)
        _rect(s, x, Inches(2.45), colw, Inches(4.0), B.STONE)
        tb = s.shapes.add_textbox(x + Inches(0.25), Inches(2.7), colw - Inches(0.5), Inches(3.6))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for j, t in enumerate(items):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.space_after = Pt(8)
            _style_run(p.add_run(), 16, False, B.INK)
            p.runs[0].text = t
            _bullet(p, 0)
    return s


def add_kpi(prs, title, items, note=""):
    """items: list[(giá trị, nhãn)] tối đa 4."""
    s = _blank(prs)
    _chrome(prs, s, title)
    n = min(len(items), 4)
    gap = Inches(0.35)
    w = (SW - 2 * MARGIN - gap * (n - 1)) / n
    for i, (val, label) in enumerate(items[:n]):
        x = MARGIN + i * (w + gap)
        dark = i == 0
        _rect(s, x, Inches(2.1), w, Inches(3.0), B.RED if dark else B.STONE)
        _rect(s, x, Inches(2.1), w, Inches(0.08), B.GOLD)
        _text(s, x + Inches(0.25), Inches(2.6), w - Inches(0.5), Inches(1.4), val, 38, True,
              B.WHITE if dark else B.RED, anchor=MSO_ANCHOR.MIDDLE)
        _text(s, x + Inches(0.25), Inches(4.0), w - Inches(0.5), Inches(0.9), label, 15, False,
              B.WHITE if dark else B.INK)
    if note:
        _text(s, MARGIN, Inches(5.5), SW - 2 * MARGIN, Inches(0.8), note, 13, False, B.GREY,
              italic=True)
    return s


def add_table(prs, title, header, rows):
    s = _blank(prs)
    _chrome(prs, s, title)
    nr, nc = len(rows) + 1, len(header)
    height = Inches(0.5) * nr
    gf = s.shapes.add_table(nr, nc, MARGIN, Inches(1.9), SW - 2 * MARGIN, height)
    tbl = gf.table
    # bỏ style mặc định (viền/màu xanh của Office)
    tblPr = tbl._tbl.tblPr
    for a in ("bandRow", "firstRow"):
        tblPr.set(a, "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is not None:
        tblPr.remove(sid)
    for r in range(nr):
        for c in range(nc):
            cell = tbl.cell(r, c)
            txt = header[c] if r == 0 else rows[r - 1][c]
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(B.RED if r == 0 else (B.STONE if r % 2 == 0 else B.WHITE))
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            _style_run(tf.paragraphs[0].add_run(), 14, r == 0, B.WHITE if r == 0 else B.INK)
            tf.paragraphs[0].runs[0].text = str(txt)
            if c > 0:
                tf.paragraphs[0].alignment = PP_ALIGN.RIGHT
    return s


def add_closing(prs, text="Xin cảm ơn", contact=""):
    s = _blank(prs)
    _rect(s, 0, 0, SW, SH, B.RED)
    _picture(s, B.star_path("color"), Inches(8.6), Inches(1.4), Inches(5.2))
    _picture(s, B.logo_path("white"), MARGIN, Inches(0.7), Inches(2.4))
    _text(s, MARGIN, Inches(3.0), Inches(7.5), Inches(1.2), text, 44, True, B.WHITE,
          anchor=MSO_ANCHOR.BOTTOM)
    _rect(s, MARGIN, Inches(4.4), Inches(1.2), Inches(0.06), B.GOLD)
    if contact:
        _text(s, MARGIN, Inches(4.6), Inches(7.5), Inches(1.0), contact, 14, False, B.WHITE)
    return s
