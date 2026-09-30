#!/usr/bin/env python3
"""Kiểm tra tuân thủ nhận diện MBV: màu ngoài bảng, font lạ, ảnh bị kéo méo.

    python check_brand.py file1.pptx file2.docx page.html style.css report.pdf
Thoát mã 1 nếu có lỗi (dùng được trong CI / hook).
"""
import os
import re
import sys
import zipfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import mbv_brand as B

ALLOWED_FONTS = {B.FONT.lower(), "arial", "helvetica", "sans-serif", "liberation sans narrow",
                 "inherit", "system-ui", "-apple-system"}
TOL = 8  # sai số mỗi kênh màu được chấp nhận
HEX = re.compile(r"#?([0-9A-Fa-f]{6})\b")


def _near(h):
    r, g, b = B.hex_rgb(h)
    return any(max(abs(r - x), abs(g - y), abs(b - z)) <= TOL
               for x, y, z in (B.hex_rgb(v) for v in B.PALETTE.values()))


def _font_ok(name):
    n = name.strip().strip("'\"").lower()
    if n.startswith("var("):
        return True
    return n in ALLOWED_FONTS or n.startswith("liberationsans") or n.startswith("arial")


def check_ooxml(path, issues):
    kind = path.lower().rsplit(".", 1)[-1]
    z = zipfile.ZipFile(path)
    parts = [n for n in z.namelist() if n.endswith(".xml") and (
        (kind == "pptx" and n.startswith("ppt/slides/slide")) or
        (kind == "docx" and re.match(r"word/(document|header\d*|footer\d*)\.xml", n)))]
    colors, fonts = {}, {}
    for n in parts:
        x = z.read(n).decode("utf8")
        rx = (r'<a:srgbClr val="([0-9A-Fa-f]{6})"' if kind == "pptx" else
              r'(?:<w:color w:val|w:fill|w:color)="([0-9A-Fa-f]{6})"')
        for m in re.findall(rx, x):
            colors.setdefault(m.upper(), set()).add(n)
        for m in re.findall(r'(?:typeface|w:ascii|w:hAnsi)="([^"]+)"', x):
            fonts.setdefault(m, set()).add(n)
        if kind == "docx" and re.search(r'w:(?:ascii|hAnsi)Theme=', x):
            issues.append((path, "font theme (không phải Liberation Sans) trong " + n))
    if kind == "docx" and "word/styles.xml" in z.namelist():
        used = set()
        for n in parts:
            used |= set(re.findall(r'<w:(?:pStyle|rStyle) w:val="([^"]+)"', z.read(n).decode("utf8")))
        sx = z.read("word/styles.xml").decode("utf8")
        for sid in used:
            m = re.search(r'<w:style [^>]*w:styleId="%s".*?</w:style>' % re.escape(sid), sx, re.S)
            if m and re.search(r'w:(?:ascii|hAnsi)Theme=', m.group(0)):
                issues.append((path, "style '%s' dùng font theme" % sid))
            for f in re.findall(r'w:ascii="([^"]+)"', m.group(0)) if m else []:
                fonts.setdefault(f, set()).add("styles:" + sid)
    for c, where in colors.items():
        if not _near(c):
            issues.append((path, "màu ngoài bảng #%s (%s)" % (c, ", ".join(sorted(where))[:60])))
    for f, where in fonts.items():
        if not _font_ok(f):
            issues.append((path, "font ngoài chuẩn '%s' (%s)" % (f, ", ".join(sorted(where))[:60])))
    if kind == "pptx":
        try:
            from pptx import Presentation
            prs = Presentation(path)
            if abs(prs.slide_width / prs.slide_height - 16 / 9) > 0.01:
                issues.append((path, "slide không phải tỷ lệ 16:9"))
            for i, s in enumerate(prs.slides, 1):
                for sh in s.shapes:
                    if sh.shape_type == 13:  # ảnh
                        iw, ih = sh.image.size
                        cw = 1 - sh.crop_left - sh.crop_right
                        chh = 1 - sh.crop_top - sh.crop_bottom
                        if abs((sh.width / sh.height) / ((iw * cw) / (ih * chh)) - 1) > 0.03:
                            issues.append((path, "slide %d: ảnh '%s' bị kéo méo" % (i, sh.name)))
        except ImportError:
            print("  (bỏ qua kiểm tra ảnh: thiếu python-pptx)")


def check_text(path, issues):
    t = open(path, encoding="utf8", errors="ignore").read()
    for h in set(HEX.findall(re.sub(r"&#?\w+;", "", t))) if "#" in t else []:
        if re.search(r"#" + h + r"\b", t) and not _near(h):
            issues.append((path, "màu ngoài bảng #%s" % h.upper()))
    for m in re.findall(r"font-family\s*:\s*([^;}\"]+)", t, re.I):
        for f in m.split(","):
            if f.strip() and not _font_ok(f):
                issues.append((path, "font ngoài chuẩn '%s'" % f.strip()))


def check_pdf(path, issues):
    try:
        import pymupdf
    except ImportError:
        print("  (bỏ qua PDF: pip install pymupdf)")
        return
    doc = pymupdf.open(path)
    fonts, colors = set(), set()
    for pg in doc:
        for b in pg.get_text("dict")["blocks"]:
            for l in b.get("lines", []):
                for s in l["spans"]:
                    fonts.add(s["font"])
                    colors.add("%06X" % s["color"])
        for d in pg.get_drawings():
            for k in ("fill", "color"):
                if d.get(k):
                    colors.add("%02X%02X%02X" % tuple(round(v * 255) for v in d[k][:3]))
    for f in fonts:
        if not _font_ok(f.split("+")[-1]) and "liberation" not in f.lower():
            issues.append((path, "font ngoài chuẩn '%s'" % f))
    for c in colors:
        if not _near(c):
            issues.append((path, "màu ngoài bảng #%s" % c))


def main(argv):
    if not argv:
        print(__doc__)
        return 2
    issues = []
    for p in argv:
        ext = p.lower().rsplit(".", 1)[-1]
        print("• kiểm tra", p)
        if ext in ("pptx", "docx"):
            check_ooxml(p, issues)
        elif ext in ("html", "htm", "css", "svg"):
            check_text(p, issues)
        elif ext == "pdf":
            check_pdf(p, issues)
        else:
            print("  (định dạng .%s chưa hỗ trợ)" % ext)
    if issues:
        print("\n✗ %d vấn đề:" % len(issues))
        for f, m in sorted(set(issues)):
            print("  - %s: %s" % (os.path.basename(f), m))
        return 1
    print("\n✓ Đạt chuẩn màu/font MBV")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
