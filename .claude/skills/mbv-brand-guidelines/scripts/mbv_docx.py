"""Dựng báo cáo A4 dọc chuẩn MBV bằng python-docx.

    import sys; sys.path.insert(0, "<skill>/scripts")
    import mbv_docx as D
    doc = D.new_report()
    D.add_cover(doc, "Báo cáo kết quả kinh doanh", "Quý III/2026", "Khối Bán lẻ · 30/09/2026")
    D.h1(doc, "1. Tổng quan"); D.para(doc, "Nội dung...")
    D.bullets(doc, ["Ý 1", "Ý 2"])
    D.table(doc, ["Chỉ tiêu", "Q2", "Q3"], [["A", "1", "2"]])
    doc.save("out.docx")
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import mbv_brand as B
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FOOTER_TEXT = "MBV · Ngân hàng TNHH MTV Việt Nam Hiện Đại"


def _rgb(h):
    return RGBColor.from_string(h.upper())


def _set_fonts(rpr_owner_rpr):
    """Đặt Liberation Sans cho mọi nhóm ký tự và xoá font theme."""
    rf = rpr_owner_rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr_owner_rpr.insert(0, rf)
    for a in list(rf.attrib):
        del rf.attrib[a]
    for a in ("ascii", "hAnsi", "eastAsia", "cs"):
        rf.set(qn("w:" + a), B.FONT)


def _style_font(style, size, bold=False, color=B.INK):
    style.font.name = B.FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = _rgb(color)
    _set_fonts(style.element.get_or_add_rPr())


def _run(p, text, size=None, bold=None, color=None, italic=None):
    r = p.add_run(text)
    r.font.name = B.FONT
    _set_fonts(r._r.get_or_add_rPr())
    if size:
        r.font.size = Pt(size)
    if bold is not None:
        r.font.bold = bold
    if italic is not None:
        r.font.italic = italic
    if color:
        r.font.color.rgb = _rgb(color)
    return r


def _shade(cell_or_par_pr_owner, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    return shd


def _cell_shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    tcpr.append(_shade(cell, fill))


def _cell_borders(cell, color=None):
    tcpr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement("w:" + side)
        if color and side == "bottom":
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:color"), color)
        else:
            e.set(qn("w:val"), "nil")
        b.append(e)
    tcpr.append(b)


def _para_border(p, side, color, sz=12, space=4):
    ppr = p._p.get_or_add_pPr()
    pb = OxmlElement("w:pBdr")
    e = OxmlElement("w:" + side)
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), str(sz))
    e.set(qn("w:space"), str(space))
    e.set(qn("w:color"), color)
    pb.append(e)
    ppr.append(pb)


def _field(p, instr):
    for kind in ("begin", None, "end"):
        r = p.add_run()
        r.font.name = B.FONT
        _set_fonts(r._r.get_or_add_rPr())
        r.font.size = Pt(9)
        r.font.color.rgb = _rgb(B.GREY)
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r._r.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = instr
            r._r.append(it)


def new_report():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2.4), Cm(2.2)
    sec.header_distance = sec.footer_distance = Cm(1.0)

    st = doc.styles
    _style_font(st["Normal"], 11)
    st["Normal"].paragraph_format.space_after = Pt(6)
    st["Normal"].paragraph_format.line_spacing = 1.4
    for name, size, color in (("Heading 1", 20, B.RED), ("Heading 2", 14, B.INK),
                              ("Heading 3", 12, B.GREY), ("Title", 30, B.RED)):
        _style_font(st[name], size, True, color)
        st[name].paragraph_format.space_before = Pt(16 if name != "Title" else 0)
        st[name].paragraph_format.space_after = Pt(6)
        st[name].paragraph_format.keep_with_next = True
    for name in ("List Bullet", "Caption"):
        _style_font(st[name], 11 if name == "List Bullet" else 9,
                    False, B.INK if name == "List Bullet" else B.GREY)
    # Tiêu đề Title mặc định có viền dưới xanh -> bỏ
    tppr = st["Title"].element.find(qn("w:pPr"))
    if tppr is not None:
        pb = tppr.find(qn("w:pBdr"))
        if pb is not None:
            tppr.remove(pb)
    return doc


def _header_footer(section):
    section.different_first_page_header_footer = True
    hp = section.header.paragraphs[0]
    hp.text = ""
    hp.add_run().add_picture(B.logo_path("color"), width=Cm(2.6))
    _para_border(hp, "bottom", B.RED, 12, 6)
    fp = section.footer.paragraphs[0]
    fp.text = ""
    _run(fp, FOOTER_TEXT + "\tTrang ", 9, False, B.GREY)
    _field(fp, "PAGE")
    fp.paragraph_format.tab_stops.add_tab_stop(Cm(16.6), alignment=2)  # phải


def add_cover(doc, title, subtitle="", meta=""):
    """Trang bìa nền đỏ (bảng 1 ô cao ~ full trang) + ngắt trang."""
    sec = doc.sections[0]
    _header_footer(sec)
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Cm(16.6)
    _cell_borders(cell)
    _cell_shade(cell, B.RED)
    trh = OxmlElement("w:trHeight")
    trh.set(qn("w:val"), str(int(Cm(23.5).twips)))
    trh.set(qn("w:hRule"), "exact")
    tbl.rows[0]._tr.get_or_add_trPr().append(trh)
    tcpr = cell._tc.get_or_add_tcPr()
    va = OxmlElement("w:vAlign")
    va.set(qn("w:val"), "top")
    tcpr.append(va)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(36)
    p.paragraph_format.left_indent = Cm(0.8)
    p.add_run().add_picture(B.logo_path("white"), width=Cm(4.6))
    sp = cell.add_paragraph()
    sp.paragraph_format.space_before = Pt(190)
    sp.paragraph_format.left_indent = Cm(0.8)
    _run(sp, title, 34, True, B.WHITE)
    sp.paragraph_format.line_spacing = 1.1
    bar = cell.add_paragraph()
    bar.paragraph_format.left_indent = Cm(0.8)
    _run(bar, "▬▬▬", 14, True, B.GOLD)
    if subtitle:
        sp2 = cell.add_paragraph()
        sp2.paragraph_format.left_indent = Cm(0.8)
        _run(sp2, subtitle, 16, False, B.WHITE)
    if meta:
        sp3 = cell.add_paragraph()
        sp3.paragraph_format.space_before = Pt(120)
        sp3.paragraph_format.left_indent = Cm(0.8)
        _run(sp3, meta, 11, True, B.GOLD)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def h1(doc, text):
    return doc.add_heading(text, 1)


def h2(doc, text):
    return doc.add_heading(text, 2)


def h3(doc, text):
    return doc.add_heading(text, 3)


def para(doc, text, bold=False, italic=False):
    p = doc.add_paragraph()
    _run(p, text, bold=bold, italic=italic)
    return p


def bullets(doc, items):
    for t in items:
        p = doc.add_paragraph(style="List Bullet")
        _run(p, t)


def caption(doc, text):
    p = doc.add_paragraph(style="Caption")
    _run(p, text, 9, False, B.GREY, True)
    return p


def callout(doc, text, title=None):
    """Khối nhấn: nền Stone, viền trái đỏ."""
    t = doc.add_table(rows=1, cols=1)
    t.autofit = False
    c = t.cell(0, 0)
    c.width = Cm(16.6)
    tcpr = c._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for side in ("top", "bottom", "right"):
        e = OxmlElement("w:" + side)
        e.set(qn("w:val"), "nil")
        b.append(e)
    l = OxmlElement("w:left")
    l.set(qn("w:val"), "single")
    l.set(qn("w:sz"), "36")
    l.set(qn("w:color"), B.RED)
    b.append(l)
    tcpr.append(b)
    _cell_shade(c, B.STONE)
    p = c.paragraphs[0]
    if title:
        _run(p, title, 11, True, B.RED)
        p = c.add_paragraph()
    _run(p, text)
    doc.add_paragraph()


def table(doc, header, rows, col_widths_cm=None):
    """Hàng tiêu đề nền đỏ chữ trắng, hàng chẵn nền Stone, cột số căn phải."""
    t = doc.add_table(rows=len(rows) + 1, cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    total = 16.6
    widths = col_widths_cm or [total / len(header)] * len(header)
    for r in range(len(rows) + 1):
        for c in range(len(header)):
            cell = t.cell(r, c)
            cell.width = Cm(widths[c])
            txt = header[c] if r == 0 else rows[r - 1][c]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.space_before = Pt(3)
            if c > 0:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _run(p, str(txt), 10.5, r == 0, B.WHITE if r == 0 else B.INK)
            _cell_borders(cell, B.LINE)
            _cell_shade(cell, B.RED if r == 0 else (B.STONE if r % 2 == 0 else B.WHITE))
    doc.add_paragraph()
    return t
